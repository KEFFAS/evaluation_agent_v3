import pandas as pd
import os
from openpyxl.styles import Font

# =========================================================
# EEE CLEANING SCRIPT
# Kenya School of Government - Matuga
# =========================================================

# ===== USER INPUT =====
file_name = input("Enter EEE Excel file name: ").strip()

# =========================================================
# LOAD FILE
# =========================================================
# Load WITHOUT headers first
df = pd.read_excel(file_name, header=None)

print("Original shape:", df.shape)

# =========================================================
# REMOVE FIRST ROW
# =========================================================
# Usually contains export artifacts / metadata
df = df.iloc[1:].reset_index(drop=True)

# =========================================================
# SET TRUE HEADER ROW
# =========================================================
df.columns = df.iloc[0]
df = df[1:].reset_index(drop=True)

# =========================================================
# REMOVE EMPTY ROWS
# =========================================================
df = df.dropna(how='all')

# =========================================================
# REMOVE UNNAMED COLUMNS
# =========================================================
df = df.loc[
    :,
    ~df.columns.astype(str).str.contains("unnamed", case=False)
]

# =========================================================
# CLEAN COLUMN NAMES
# IMPORTANT:
# Preserve original question wording
# =========================================================
df.columns = (
    df.columns
    .astype(str)
    .str.strip()
)

# =========================================================
# CLEAN TEXT VALUES
# =========================================================
for col in df.columns:

    if df[col].dtype == "object":

        df[col] = (
            df[col]
            .astype(str)
            .str.strip()
        )

# =========================================================
# CONVERT NUMERIC COLUMNS
# Smart numeric detection
# =========================================================
for col in df.columns:

    converted = pd.to_numeric(df[col], errors='coerce')

    # If majority values are numeric -> convert
    if converted.notna().sum() > len(df) * 0.5:
        df[col] = converted

# =========================================================
# REMOVE DUPLICATES
# =========================================================
#df = df.drop_duplicates()

# =========================================================
# STANDARDIZE ONLY ESSENTIAL IDENTIFIERS
# DO NOT ALTER QUESTION HEADERS
# =========================================================
rename_map = {
    "Program Name": "Program Title",
    "Coordinator": "Coordinator Name"
}

df = df.rename(columns=rename_map)

# =========================================================
# REMOVE NON-ESSENTIAL SYSTEM COLUMNS
# =========================================================
drop_cols = [
    "Average Rating",
    "Status",
    "Course Duration (Days)"
]

existing_drop_cols = [
    col for col in drop_cols
    if col in df.columns
]

df = df.drop(columns=existing_drop_cols)

# =========================================================
# DETECT IMPORTANT SECTION COLUMNS
# Helps downstream scripts
# =========================================================

# ----- Course Objectives -----
objective_cols = [
    col for col in df.columns
    if "objective" in str(col).lower()
]

# ----- Personal Expectations -----
expectation_cols = [
    col for col in df.columns
    if "expectation" in str(col).lower()
]

# ----- Institution Comparison -----
comparison_cols = [
    col for col in df.columns
    if "similar institution" in str(col).lower()
]

# ----- Qualitative Columns -----
qualitative_cols = [
    col for col in df.columns
    if any(keyword in str(col).lower() for keyword in [
        "suggest",
        "comment",
        "interest",
        "area",
        "training programs"
    ])
]

# =========================================================
# PRINT DETECTED STRUCTURE
# =========================================================
print("\n==============================")
print("Detected Key Sections")
print("==============================")

print("\nCourse Objective Columns:")
for col in objective_cols:
    print("-", col)

print("\nExpectation Columns:")
for col in expectation_cols:
    print("-", col)

print("\nInstitution Comparison Columns:")
for col in comparison_cols:
    print("-", col)

print("\nQualitative Columns:")
for col in qualitative_cols:
    print("-", col)

# =========================================================
# SAVE CLEANED FILE
# =========================================================
base_name = os.path.splitext(file_name)[0]

output_file = f"{base_name}_eee_cleaned.xlsx"

# Prevent overwrite
if os.path.exists(output_file):
    output_file = f"{base_name}_eee_cleaned_new.xlsx"

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:

    df.to_excel(
        writer,
        index=False,
        sheet_name='Cleaned Data'
    )

    worksheet = writer.sheets['Cleaned Data']

    # Bold headers
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

print("\n===================================")
print("✅ Cleaned EEE file saved as:")
print(output_file)
print("===================================")


import pandas as pd
import os
from collections import Counter
from openpyxl.styles import Font

# =========================================================
# EEE ANALYSIS SCRIPT
# Kenya School of Government - Matuga
# =========================================================

# ===== INPUT =====
file_name = input("Enter cleaned EEE file name: ").strip()

# ===== LOAD FILE =====
df = pd.read_excel(file_name)

print("Loaded:", file_name)
print("Shape:", df.shape)

# =========================================================
# BASIC DETAILS
# =========================================================
program_title = (
    df["Program Title"].iloc[0]
    if "Program Title" in df.columns
    else "N/A"
)

coordinator = (
    df["Coordinator Name"].iloc[0]
    if "Coordinator Name" in df.columns
    else "N/A"
)

program_code = (
    df["Program Code"].iloc[0]
    if "Program Code" in df.columns
    else "N/A"
)

venue = (
    df["Venue / Campus"].iloc[0]
    if "Venue / Campus" in df.columns
    else "N/A"
)

assistant = (
    df["Program Assistant Name"].iloc[0]
    if "Program Assistant Name" in df.columns
    else "N/A"
)

# =========================================================
# DETECT NUMERIC RATING COLUMNS
# =========================================================
rating_cols = [
    col for col in df.columns
    if df[col].dtype in ["int64", "float64", "Int64"]
]

# Remove non-rating numeric columns
exclude = [
    "Timetable No"
]

rating_cols = [
    col for col in rating_cols
    if col not in exclude
]

# =========================================================
# DETECT SECTION COLUMNS
# =========================================================

# ----- Objectives -----
objective_col = next(
    (
        col for col in rating_cols
        if "objective" in col.lower()
    ),
    None
)

# ----- Expectations -----
expectation_col = next(
    (
        col for col in rating_cols
        if "expectation" in col.lower()
    ),
    None
)

# ----- Institution Comparison -----
comparison_col = next(
    (
        col for col in rating_cols
        if "similar institution" in col.lower()
    ),
    None
)

# =========================================================
# SECTION 1 ANALYSIS
# Course Objectives Achievement
# =========================================================
section1 = []

if objective_col:

    counts = df[objective_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "Excellent",
        4: "Very Good",
        3: "Satisfactory",
        2: "Poor",
        1: "Very Poor"
    }

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        section1.append([
            labels[score],
            pct
        ])

section1_df = pd.DataFrame(
    section1,
    columns=["Rating", "Percentage of Respondents"]
)

# =========================================================
# SECTION 2 ANALYSIS
# Personal Expectations
# =========================================================
section2 = []

if expectation_col:

    counts = df[expectation_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "5 - Great Extent",
        4: "4 - Some Extent",
        3: "3 - Satisfactory",
        2: "2 - Not Sure",
        1: "1 - Not at All"
    }

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        section2.append([
            labels[score],
            pct
        ])

section2_df = pd.DataFrame(
    section2,
    columns=["Rating", "Percentage of Respondents"]
)

# =========================================================
# SECTION 3 ANALYSIS
# Specific Aspects Table
# =========================================================
specific_aspects = []

for col in rating_cols:

    if col not in [objective_col, expectation_col, comparison_col]:

        counts = df[col].value_counts().to_dict()
        total = sum(counts.values())

        row = [
            col,

            round((counts.get(5,0)/total)*100,1)
            if total > 0 else 0,

            round((counts.get(4,0)/total)*100,1)
            if total > 0 else 0,

            round((counts.get(3,0)/total)*100,1)
            if total > 0 else 0,

            round((counts.get(2,0)/total)*100,1)
            if total > 0 else 0,

            round((counts.get(1,0)/total)*100,1)
            if total > 0 else 0
        ]

        specific_aspects.append(row)

section3_df = pd.DataFrame(
    specific_aspects,
    columns=[
        "ASPECT OF THE PROGRAM",
        "Excellent %",
        "Very Good %",
        "Satisfactory %",
        "Poor %",
        "Very Poor %"
    ]
)

# =========================================================
# SECTION 8 ANALYSIS
# Institution Comparison
# =========================================================
section8 = []

if comparison_col:

    counts = df[comparison_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "5 - Very High",
        4: "4 - High",
        3: "3 - Average",
        2: "2 - Low",
        1: "1 - Very Low"
    }

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        section8.append([
            labels[score],
            pct
        ])

section8_df = pd.DataFrame(
    section8,
    columns=["Rating", "Percentage of Respondents"]
)

# =========================================================
# QUALITATIVE SECTION MAPPING
# STRICTLY BASED ON KSG STRUCTURE
# =========================================================

qualitative_mapping = {
    "Suggestions": None,
    "Areas to Add": None,
    "Interest in Other KSG Programmes": None,
    "Additional Training Areas": None,
    "General Comments": None
}

for col in df.columns:

    col_lower = str(col).lower()

    if "suggestions on aspects" in col_lower:
        qualitative_mapping["Suggestions"] = col

    elif "other areas you would like added" in col_lower:
        qualitative_mapping["Areas to Add"] = col

    elif "other ksg training programs" in col_lower:
        qualitative_mapping["Interest in Other KSG Programmes"] = col

    elif "other training programs not currently offered" in col_lower:
        qualitative_mapping["Additional Training Areas"] = col

    elif "other comments" in col_lower:
        qualitative_mapping["General Comments"] = col

# =========================================================
# EXTRACT QUALITATIVE TEXT
# =========================================================
qualitative_outputs = {}

for section, column in qualitative_mapping.items():

    if column and column in df.columns:

        text = " ".join(
            df[column]
            .dropna()
            .astype(str)
        )

        qualitative_outputs[section] = text

    else:
        qualitative_outputs[section] = ""

# =========================================================
# SAVE ANALYSIS
# =========================================================
base_name = os.path.splitext(file_name)[0]

output_file = f"{base_name}_eee_analysis.xlsx"

if os.path.exists(output_file):
    output_file = f"{base_name}_eee_analysis_new.xlsx"

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:

    # =====================================================
    # DETAILS SHEET
    # =====================================================
    details_df = pd.DataFrame({
        "Field": [
            "Program Title",
            "Coordinator",
            "Program Code",
            "Venue",
            "Program Assistant"
        ],
        "Value": [
            program_title,
            coordinator,
            program_code,
            venue,
            assistant
        ]
    })

    details_df.to_excel(
        writer,
        sheet_name="Program Details",
        index=False
    )

    # =====================================================
    # SECTION 1
    # =====================================================
    section1_df.to_excel(
        writer,
        sheet_name="Section 1 Objectives",
        index=False
    )

    # =====================================================
    # SECTION 2
    # =====================================================
    section2_df.to_excel(
        writer,
        sheet_name="Section 2 Expectations",
        index=False
    )

    # =====================================================
    # SECTION 3
    # =====================================================
    section3_df.to_excel(
        writer,
        sheet_name="Section 3 Specific Aspects",
        index=False
    )

    # =====================================================
    # SECTION 8
    # =====================================================
    section8_df.to_excel(
        writer,
        sheet_name="Section 8 Comparison",
        index=False
    )

    # =====================================================
    # QUALITATIVE SHEET
    # =====================================================
    qualitative_sheet = []

    for section, text in qualitative_outputs.items():

        qualitative_sheet.append([
            section,
            text
        ])

    qualitative_df = pd.DataFrame(
        qualitative_sheet,
        columns=["Section", "Responses"]
    )

    qualitative_df.to_excel(
        writer,
        sheet_name="Qualitative Responses",
        index=False
    )

    # =====================================================
    # BOLD HEADERS
    # =====================================================
    for sheet in writer.sheets.values():

        for cell in sheet[1]:
            cell.font = Font(bold=True)

print("\n===================================")
print("✅ EEE analysis completed")
print("Saved as:")
print(output_file)
print("===================================")


import pandas as pd
import os
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# =========================================================
# INTEGRATED EEE REPORT GENERATOR
# Kenya School of Government - Matuga
# =========================================================

# =========================================================
# TABLE BORDER FUNCTION
# =========================================================
def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement('w:tblBorders')

    for border_name in [
        'top',
        'left',
        'bottom',
        'right',
        'insideH',
        'insideV'
    ]:

        border = OxmlElement(f'w:{border_name}')

        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:color'), '000000')

        borders.append(border)

    tblPr.append(borders)

# =========================================================
# INPUT FILE
# =========================================================
file_name = input("Enter cleaned EEE file: ").strip()

# =========================================================
# LOAD FILE
# =========================================================
df = pd.read_excel(file_name)

print("Loaded:", file_name)
print("Shape:", df.shape)

# =========================================================
# PROGRAM DETAILS
# =========================================================
program_title = (
    df["Program Title"].iloc[0]
    if "Program Title" in df.columns
    else "N/A"
)

coordinator = (
    df["Coordinator Name"].iloc[0]
    if "Coordinator Name" in df.columns
    else "N/A"
)

program_code = (
    df["Program Code"].iloc[0]
    if "Program Code" in df.columns
    else "N/A"
)

venue = (
    df["Venue / Campus"].iloc[0]
    if "Venue / Campus" in df.columns
    else "N/A"
)

assistant = (
    df["Program Assistant Name"].iloc[0]
    if "Program Assistant Name" in df.columns
    else "N/A"
)

duration = input("Enter Program Duration: ")

# =========================================================
# DETECT NUMERIC RATING COLUMNS
# =========================================================
rating_cols = [
    col for col in df.columns
    if df[col].dtype in ["int64", "float64", "Int64"]
]

exclude = [
    "Timetable No"
]

rating_cols = [
    col for col in rating_cols
    if col not in exclude
]

# =========================================================
# DETECT KEY SECTION COLUMNS
# =========================================================

objective_col = next(
    (
        col for col in rating_cols
        if "objective" in col.lower()
    ),
    None
)

expectation_col = next(
    (
        col for col in rating_cols
        if "expectation" in col.lower()
    ),
    None
)

comparison_col = next(
    (
        col for col in rating_cols
        if "similar institution" in col.lower()
    ),
    None
)

# =========================================================
# CREATE DOCUMENT
# =========================================================
doc = Document()

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

# =========================================================
# HEADER
# =========================================================
doc.add_paragraph("KSG/17/EOEEF/07")
doc.add_paragraph("KENYA SCHOOL OF GOVERNMENT")
doc.add_paragraph("MATUGA")
doc.add_heading("END-OF-EVENT EVALUATION FORM", level=1)

# =========================================================
# PROGRAM DETAILS TABLE
# =========================================================
table = doc.add_table(rows=3, cols=4)

table.cell(0, 0).text = "PROGRAMME TITLE:"
table.cell(0, 1).text = str(program_title)

table.cell(0, 2).text = "DURATION:"
table.cell(0, 3).text = str(duration)

table.cell(1, 0).text = "PROGRAM CODE:"
table.cell(1, 1).text = str(program_code)

table.cell(1, 2).text = "VENUE:"
table.cell(1, 3).text = str(venue)

table.cell(2, 0).text = "COORDINATOR:"
table.cell(2, 1).text = str(coordinator)

table.cell(2, 2).text = "PROGRAM ASST:"
table.cell(2, 3).text = str(assistant)

set_table_borders(table)

# =========================================================
# INTRODUCTION
# =========================================================
doc.add_heading("A. PROGRAMME EVALUATION", level=2)

doc.add_paragraph(
    "KSG is committed to providing quality-training programmes "
    "to its customers. We therefore request you to complete "
    "this evaluation form candidly as you can in order to help "
    "us fulfil our commitment to continuously improve our programmes."
)

# =========================================================
# SECTION 1
# =========================================================
doc.add_heading("1. Course Objectives Achievement", level=2)

table1 = doc.add_table(rows=1, cols=2)

table1.rows[0].cells[0].text = "Rating"
table1.rows[0].cells[1].text = "Percentage of Respondents"

if objective_col:

    counts = df[objective_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "Excellent",
        4: "Very Good",
        3: "Satisfactory",
        2: "Poor",
        1: "Very Poor"
    }

    for score in [5,4,3,2,1]:

        row = table1.add_row().cells

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        row[0].text = labels[score]
        row[1].text = str(pct)

set_table_borders(table1)

# =========================================================
# SECTION 2
# =========================================================
doc.add_heading(
    "2. Fulfilment of personal expectations",
    level=2
)

table2 = doc.add_table(rows=1, cols=2)

table2.rows[0].cells[0].text = "Rating"
table2.rows[0].cells[1].text = "Percentage of Respondents"

if expectation_col:

    counts = df[expectation_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "5 - Great Extent",
        4: "4 - Some Extent",
        3: "3 - Satisfactory",
        2: "2 - Not Sure",
        1: "1 - Not at All"
    }

    for score in [5,4,3,2,1]:

        row = table2.add_row().cells

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        row[0].text = labels[score]
        row[1].text = str(pct)

set_table_borders(table2)

# =========================================================
# SECTION 3
# =========================================================
doc.add_heading(
    "3. Ratings on specific aspects of the training program",
    level=2
)

table3 = doc.add_table(rows=1, cols=6)

headers = [
    "ASPECT OF THE PROGRAM",
    "Excellent %",
    "Very Good %",
    "Satisfactory %",
    "Poor %",
    "Very poor %"
]

for i, h in enumerate(headers):
    table3.rows[0].cells[i].text = h

for col in rating_cols:

    if col not in [
        objective_col,
        expectation_col,
        comparison_col
    ]:

        counts = df[col].value_counts().to_dict()
        total = sum(counts.values())

        row = table3.add_row().cells

        row[0].text = str(col)

        row[1].text = str(
            round((counts.get(5,0)/total)*100,1)
            if total > 0 else 0
        )

        row[2].text = str(
            round((counts.get(4,0)/total)*100,1)
            if total > 0 else 0
        )

        row[3].text = str(
            round((counts.get(3,0)/total)*100,1)
            if total > 0 else 0
        )

        row[4].text = str(
            round((counts.get(2,0)/total)*100,1)
            if total > 0 else 0
        )

        row[5].text = str(
            round((counts.get(1,0)/total)*100,1)
            if total > 0 else 0
        )

set_table_borders(table3)

# =========================================================
# QUALITATIVE SECTION MAPPING
# =========================================================
qualitative_mapping = {
    "4. Suggestions on the aspects listed in (3) above.":
        "suggestions on aspects",

    "5. Areas to be added to this training programme":
        "other areas you would like added",

    "6. Interest in other KSG programmes":
        "other ksg training programs",

    "7. Interest in additional training areas not currently offered by KSG":
        "other training programs not currently offered",

    "9. General Comments":
        "other comments"
}

# =========================================================
# QUALITATIVE OUTPUT (PARAGRAPH STYLE)
# =========================================================
for section_title, keyword in qualitative_mapping.items():

    matching_cols = [
        col for col in df.columns
        if keyword in str(col).lower()
    ]

    if matching_cols:

        col = matching_cols[0]

        # Join responses into flowing paragraph
        responses = (
            df[col]
            .dropna()
            .astype(str)
            .str.strip()
        )

        # Remove blanks and duplicates
        responses = responses[
            responses != ""
        ].drop_duplicates()

        text = " ".join(responses)

        doc.add_heading(section_title, level=2)

        doc.add_paragraph(text)

# =========================================================
# SECTION 8
# =========================================================
doc.add_heading(
    "8. Rating of KSG’s Training Compared to Similar Institutions",
    level=2
)

table8 = doc.add_table(rows=1, cols=2)

table8.rows[0].cells[0].text = "Rating"
table8.rows[0].cells[1].text = "Percentage of Respondents"

if comparison_col:

    counts = df[comparison_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "5 - Very High",
        4: "4 - High",
        3: "3 - Average",
        2: "2 - Low",
        1: "1 - Very Low"
    }

    for score in [5,4,3,2,1]:

        row = table8.add_row().cells

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        row[0].text = labels[score]
        row[1].text = str(pct)

set_table_borders(table8)

# =========================================================
# SECTION 10
# =========================================================
doc.add_heading(
    "10. Key Recommendations",
    level=2
)

doc.add_paragraph(
    "Recommendations to be added after qualitative analysis."
)

# =========================================================
# SIGNATURE SECTION
# =========================================================
doc.add_paragraph(
    "\nPrepared by………..………..….………...…………"
    "Date…………….…………...Signature……………"
)

doc.add_paragraph(
    "Confirmed by………..….………...…………………"
    "Date…………………………Signature……………"
)

doc.add_paragraph(
    "Approved by………………...…..…….………..……"
    "Date…………………………Signature……………"
)

# =========================================================
# SAVE DOCUMENT
# =========================================================
base_name = os.path.splitext(file_name)[0]

output_file = f"{base_name}_EEE_Report.docx"

doc.save(output_file)

print("\n===================================")
print("✅ EEE Report Generated")
print("Saved as:")
print(output_file)
print("===================================")


import pandas as pd
import os
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openai import OpenAI
from dotenv import load_dotenv

# =========================================================
# LOAD ENVIRONMENT VARIABLES
# =========================================================
load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)

# =========================================================
# TABLE BORDER FUNCTION
# =========================================================
def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement('w:tblBorders')

    for border_name in [
        'top',
        'left',
        'bottom',
        'right',
        'insideH',
        'insideV'
    ]:

        border = OxmlElement(f'w:{border_name}')

        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:color'), '000000')

        borders.append(border)

    tblPr.append(borders)

# =========================================================
# LLM FUNCTION
# =========================================================
def generate_text(prompt):

    response = client.chat.completions.create(

        model="gpt-5-nano-2025-08-07",

        messages=[

            {
                "role": "system",
                "content":
                (
                    "You are an institutional monitoring and evaluation officer "
                    "writing formal Kenya School of Government evaluation reports. "
                    "Write in a professional, concise, evidence-based and human tone. "
                    "Avoid exaggerated language, repetition, and generic AI wording."
                )
            },

            {
                "role": "user",
                "content": prompt
            }
        ]

        
    )

    return response.choices[0].message.content.strip()

# =========================================================
# INPUT FILE
# =========================================================
file_name = input("Enter cleaned EEE file: ").strip()

# =========================================================
# LOAD FILE
# =========================================================
df = pd.read_excel(file_name)

print("Loaded:", file_name)

# =========================================================
# PROGRAM DETAILS
# =========================================================
program_title = (
    df["Program Title"].iloc[0]
    if "Program Title" in df.columns
    else "N/A"
)

coordinator = (
    df["Coordinator Name"].iloc[0]
    if "Coordinator Name" in df.columns
    else "N/A"
)

program_code = (
    df["Program Code"].iloc[0]
    if "Program Code" in df.columns
    else "N/A"
)

venue = (
    df["Venue / Campus"].iloc[0]
    if "Venue / Campus" in df.columns
    else "N/A"
)

assistant = (
    df["Program Assistant Name"].iloc[0]
    if "Program Assistant Name" in df.columns
    else "N/A"
)

duration = input("Enter Program Duration: ")

# =========================================================
# DETECT NUMERIC COLUMNS
# =========================================================
rating_cols = [
    col for col in df.columns
    if df[col].dtype in ["int64", "float64", "Int64"]
]

exclude = ["Timetable No"]

rating_cols = [
    col for col in rating_cols
    if col not in exclude
]

# =========================================================
# DETECT SPECIAL COLUMNS
# =========================================================
objective_col = next(
    (
        col for col in rating_cols
        if "objective" in col.lower()
    ),
    None
)

expectation_col = next(
    (
        col for col in rating_cols
        if "expectation" in col.lower()
    ),
    None
)

comparison_col = next(
    (
        col for col in rating_cols
        if "similar institution" in col.lower()
    ),
    None
)

# =========================================================
# CREATE DOCUMENT
# =========================================================
doc = Document()

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

# =========================================================
# HEADER
# =========================================================
doc.add_paragraph("KSG/17/EOEEF/07")
doc.add_paragraph("KENYA SCHOOL OF GOVERNMENT")
doc.add_paragraph("MATUGA")
doc.add_heading("END-OF-EVENT EVALUATION REPORT", level=1)

# =========================================================
# PROGRAM DETAILS TABLE
# =========================================================
table = doc.add_table(rows=3, cols=4)

table.cell(0, 0).text = "PROGRAMME TITLE:"
table.cell(0, 1).text = str(program_title)

table.cell(0, 2).text = "DURATION:"
table.cell(0, 3).text = str(duration)

table.cell(1, 0).text = "PROGRAM CODE:"
table.cell(1, 1).text = str(program_code)

table.cell(1, 2).text = "VENUE:"
table.cell(1, 3).text = str(venue)

table.cell(2, 0).text = "COORDINATOR:"
table.cell(2, 1).text = str(coordinator)

table.cell(2, 2).text = "PROGRAM ASST:"
table.cell(2, 3).text = str(assistant)

set_table_borders(table)

# =========================================================
# INTRODUCTION
# =========================================================
doc.add_heading("A. PROGRAMME EVALUATION", level=2)

doc.add_paragraph(
    "KSG is committed to providing quality training programmes "
    "to its customers. The evaluation findings presented in this "
    "report are intended to support continuous improvement in "
    "programme design, delivery and participant experience."
)

# =========================================================
# SECTION 1
# =========================================================
doc.add_heading(
    "1. Course Objectives Achievement",
    level=2
)

table1 = doc.add_table(rows=1, cols=2)

table1.rows[0].cells[0].text = "Rating"
table1.rows[0].cells[1].text = "Percentage of Respondents"

section1_results = []

if objective_col:

    counts = df[objective_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "Excellent",
        4: "Very Good",
        3: "Satisfactory",
        2: "Poor",
        1: "Very Poor"
    }

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        row = table1.add_row().cells

        row[0].text = labels[score]
        row[1].text = str(pct)

        section1_results.append(
            f"{labels[score]} = {pct}%"
        )

set_table_borders(table1)

prompt = f"""
Interpret the following course objective achievement results:

{section1_results}

Write one concise institutional paragraph similar to a Kenya School of Government evaluation report.
"""

doc.add_paragraph(generate_text(prompt))

# =========================================================
# SECTION 2
# =========================================================
doc.add_heading(
    "2. Fulfilment of Personal Expectations",
    level=2
)

table2 = doc.add_table(rows=1, cols=2)

table2.rows[0].cells[0].text = "Rating"
table2.rows[0].cells[1].text = "Percentage of Respondents"

section2_results = []

if expectation_col:

    counts = df[expectation_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "Great Extent",
        4: "Some Extent",
        3: "Satisfactory",
        2: "Not Sure",
        1: "Not at All"
    }

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        row = table2.add_row().cells

        row[0].text = labels[score]
        row[1].text = str(pct)

        section2_results.append(
            f"{labels[score]} = {pct}%"
        )

set_table_borders(table2)

prompt = f"""
Interpret the following participant expectation fulfilment results:

{section2_results}

Write one concise institutional paragraph similar to a Kenya School of Government evaluation report.
"""

doc.add_paragraph(generate_text(prompt))

# =========================================================
# SECTION 3
# =========================================================
doc.add_heading(
    "3. Ratings on Specific Aspects of the Training Programme",
    level=2
)

table3 = doc.add_table(rows=1, cols=6)

headers = [
    "ASPECT OF THE PROGRAMME",
    "Excellent %",
    "Very Good %",
    "Satisfactory %",
    "Poor %",
    "Very Poor %"
]

for i, h in enumerate(headers):
    table3.rows[0].cells[i].text = h

section3_summary = []

for col in rating_cols:

    if col not in [
        objective_col,
        expectation_col,
        comparison_col
    ]:

        counts = df[col].value_counts().to_dict()
        total = sum(counts.values())

        excellent = round((counts.get(5,0)/total)*100,1) if total else 0
        very_good = round((counts.get(4,0)/total)*100,1) if total else 0
        satisfactory = round((counts.get(3,0)/total)*100,1) if total else 0
        poor = round((counts.get(2,0)/total)*100,1) if total else 0
        very_poor = round((counts.get(1,0)/total)*100,1) if total else 0

        row = table3.add_row().cells

        row[0].text = str(col)
        row[1].text = str(excellent)
        row[2].text = str(very_good)
        row[3].text = str(satisfactory)
        row[4].text = str(poor)
        row[5].text = str(very_poor)

        section3_summary.append(
            f"{col}: Excellent={excellent}%, Very Good={very_good}%"
        )

set_table_borders(table3)

prompt = f"""
Interpret the following programme aspect ratings:

{section3_summary}

Write one concise institutional paragraph highlighting overall participant satisfaction trends.
"""

doc.add_paragraph(generate_text(prompt))

# =========================================================
# QUALITATIVE SECTIONS
# =========================================================
qualitative_mapping = {
    "4. Suggestions on the aspects listed in (3) above.":
        "suggestions on aspects",

    "5. Areas to be added to this training programme":
        "other areas you would like added",

    "6. Interest in other KSG programmes":
        "other ksg training programs",

    "7. Interest in additional training areas not currently offered by KSG":
        "other training programs not currently offered",

    "9. General Comments":
        "other comments"
}

for section_title, keyword in qualitative_mapping.items():

    matching_cols = [
        col for col in df.columns
        if keyword in str(col).lower()
    ]

    if matching_cols:

        col = matching_cols[0]

        responses = (
            df[col]
            .dropna()
            .astype(str)
            .str.strip()
        )

        responses = responses[
            responses != ""
        ].drop_duplicates()

        joined_text = " ".join(responses)

        prompt = f"""
The following are participant responses from a Kenya School of Government training evaluation:

{joined_text}

Write one professional paragraph summarizing the most recurring themes only.
Avoid repetition and avoid listing every response individually.
"""

        doc.add_heading(section_title, level=2)

        doc.add_paragraph(
            generate_text(prompt)
        )

# =========================================================
# SECTION 8
# =========================================================
doc.add_heading(
    "8. Rating of KSG’s Training Compared to Similar Institutions",
    level=2
)

table8 = doc.add_table(rows=1, cols=2)

table8.rows[0].cells[0].text = "Rating"
table8.rows[0].cells[1].text = "Percentage of Respondents"

comparison_results = []

if comparison_col:

    counts = df[comparison_col].value_counts().to_dict()
    total = sum(counts.values())

    labels = {
        5: "Very High",
        4: "High",
        3: "Average",
        2: "Low",
        1: "Very Low"
    }

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        row = table8.add_row().cells

        row[0].text = labels[score]
        row[1].text = str(pct)

        comparison_results.append(
            f"{labels[score]} = {pct}%"
        )

set_table_borders(table8)

prompt = f"""
Interpret the following institutional comparison ratings:

{comparison_results}

Write one concise institutional paragraph similar to a Kenya School of Government evaluation report.
"""

doc.add_paragraph(generate_text(prompt))

# =========================================================
# SECTION 10
# =========================================================
doc.add_heading(
    "10. Key Recommendations",
    level=2
)

recommendation_text = ""

for col in df.columns:

    if any(keyword in str(col).lower() for keyword in [
        "suggest",
        "comment",
        "area"
    ]):

        recommendation_text += " ".join(
            df[col]
            .dropna()
            .astype(str)
        )

prompt = f"""
Based on the following participant feedback:

{recommendation_text}

Generate concise actionable institutional recommendations for improving future programmes.
"""

doc.add_paragraph(
    generate_text(prompt)
)

# =========================================================
# SIGNATURE SECTION
# =========================================================
doc.add_paragraph(
    "\nPrepared by………..………..….………...…………"
    "Date…………….…………...Signature……………"
)

doc.add_paragraph(
    "Confirmed by………..….………...…………………"
    "Date…………………………Signature……………"
)

doc.add_paragraph(
    "Approved by………………...…..…….………..……"
    "Date…………………………Signature……………"
)

# =========================================================
# SAVE DOCUMENT
# =========================================================
base_name = os.path.splitext(file_name)[0]

output_file = f"{base_name}_EEE_Report_LLM.docx"

doc.save(output_file)

print("\n===================================")
print("✅ EEE LLM REPORT GENERATED")
print("Saved as:")
print(output_file)
print("===================================")




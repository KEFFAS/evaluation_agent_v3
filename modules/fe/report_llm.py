import os
import pandas as pd

from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from openai import OpenAI
from dotenv import load_dotenv


# =========================================================
# LOAD ENVIRONMENT
# =========================================================

load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)


# =========================================================
# TABLE BORDER FUNCTION
# =========================================================

def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for border_name in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ]:

        border = OxmlElement(f"w:{border_name}")

        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")

        borders.append(border)

    tblPr.append(borders)


# =========================================================
# FORMAT CELL
# =========================================================

def format_cell(cell, bold=False, size=9):

    for paragraph in cell.paragraphs:

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for run in paragraph.runs:

            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.bold = bold


# =========================================================
# CALCULATE SESSION STATISTICS
# =========================================================

def get_session_statistics(group, rating_cols):

    results = {}

    means = []

    # Total respondents =
    # Number of rows/respondents who submitted evaluation
    total_respondents = len(group)

    for col in rating_cols:

        values = pd.to_numeric(
            group[col],
            errors="coerce"
        )

        count_5 = int((values == 5).sum())
        count_4 = int((values == 4).sum())
        count_3 = int((values == 3).sum())
        count_2 = int((values == 2).sum())
        count_1 = int((values == 1).sum())

        total_valid = (
            count_5 +
            count_4 +
            count_3 +
            count_2 +
            count_1
        )

        if total_valid > 0:

            mean = (
                (5 * count_5) +
                (4 * count_4) +
                (3 * count_3) +
                (2 * count_2) +
                (1 * count_1)
            ) / total_valid

            means.append(mean)

        else:

            mean = 0

        distribution = (
            f"{count_5}/"
            f"{count_4}/"
            f"{count_3}/"
            f"{count_2}/"
            f"{count_1}"
        )

        results[col] = {

            "distribution": distribution,

            "mean": round(mean, 2)

        }

    # Session overall mean

    if means:

        session_mean = round(
            sum(means) / len(means),
            2
        )

    else:

        session_mean = 0

    # Composite score

    composite_score = round(
        (session_mean / 5) * 100,
        1
    )

    return {

        "ratings": results,

        "session_mean": session_mean,

        "composite_score": composite_score,

        "total_respondents": total_respondents
    }


# =========================================================
# PERFORMANCE INTERPRETATION
# =========================================================

def get_performance_level(composite_score):

    if composite_score >= 90:
        return "Excellent"

    elif composite_score >= 80:
        return "Very Good"

    elif composite_score >= 70:
        return "Good"

    elif composite_score >= 60:
        return "Satisfactory"

    else:
        return "Needs Improvement"


# =========================================================
# AUTOMATED HOD COMMENT
# =========================================================

def generate_hod_comment(
    composite_score,
    strongest,
    weakest
):

    level = get_performance_level(
        composite_score
    )

    if composite_score >= 90:

        return (
            f"The facilitator demonstrated {level.lower()} "
            f"performance with an overall composite score of "
            f"{composite_score}%. Strong performance was "
            f"particularly observed in {strongest}. The "
            f"facilitator should sustain this standard while "
            f"continuing to strengthen {weakest}."
        )

    elif composite_score >= 80:

        return (
            f"The facilitator demonstrated {level.lower()} "
            f"performance with an overall composite score of "
            f"{composite_score}%. Performance was strongest "
            f"in {strongest}, while further attention to "
            f"{weakest} would enhance overall effectiveness."
        )

    elif composite_score >= 70:

        return (
            f"The facilitator demonstrated good overall "
            f"performance with a composite score of "
            f"{composite_score}%. While {strongest} was a "
            f"key strength, improvement in {weakest} would "
            f"further enhance facilitation effectiveness."
        )

    else:

        return (
            f"The facilitator achieved a composite score of "
            f"{composite_score}%, indicating a need for "
            f"improvement. Priority attention should be "
            f"given to {weakest}, while building on the "
            f"relative strength observed in {strongest}."
        )


# =========================================================
# AUTOMATED RECOMMENDATION
# =========================================================

def generate_recommendation(
    composite_score,
    weakest
):

    if composite_score >= 90:

        return (
            f"The facilitator should sustain the strong "
            f"performance demonstrated and continue refining "
            f"{weakest}. Continued deployment in similar "
            f"training programmes is recommended."
        )

    elif composite_score >= 80:

        return (
            f"The facilitator should maintain the positive "
            f"performance demonstrated while undertaking "
            f"targeted improvement in {weakest}."
        )

    elif composite_score >= 70:

        return (
            f"Targeted support and continuous professional "
            f"development should be considered, particularly "
            f"in the area of {weakest}."
        )

    else:

        return (
            f"A structured improvement plan should be "
            f"developed, with targeted support in {weakest} "
            f"and follow-up evaluation in subsequent "
            f"training assignments."
        )


# =========================================================
# STRATEGIC ADVICE
# =========================================================

def generate_strategic_advice(
    composite_score,
    strongest,
    weakest
):

    if composite_score >= 90:

        return (
            f"Management may leverage the facilitator's "
            f"strength in {strongest} through continued "
            f"deployment and peer learning, while supporting "
            f"continuous improvement in {weakest}."
        )

    elif composite_score >= 80:

        return (
            f"Strategically, targeted feedback and coaching "
            f"should focus on {weakest} while maintaining "
            f"the facilitator's strengths in {strongest}."
        )

    else:

        return (
            f"Management should consider targeted mentoring "
            f"and professional development, particularly in "
            f"{weakest}, while building on the facilitator's "
            f"strength in {strongest}."
        )


# =========================================================
# LLM QUALITATIVE ANALYSIS
# =========================================================

def analyze_qualitative(text):

    if not text.strip():

        return (
            "No qualitative comments were provided.\n\n"
            "No specific suggestions for improvement were provided."
        )

    text = text[:5000]

    prompt = f"""
The following comments were provided by participants during
evaluation of a facilitator at the Kenya School of Government.

Participant Feedback:

{text}

Write exactly TWO concise professional paragraphs.

Paragraph 1:
Summarize the most recurring positive feedback regarding
facilitation, delivery style, subject mastery, participant
engagement, communication, responsiveness and overall
teaching effectiveness.

Paragraph 2:
Summarize the most recurring suggestions for improvement.

Requirements:

- Formal institutional language
- Human tone
- Evidence based
- Do not invent feedback
- No bullet points
- No headings
- No repetition
- Suitable for an official KSG report
"""

    response = client.chat.completions.create(

        model="gpt-5-nano-2025-08-07",

        messages=[

            {
                "role": "system",

                "content": (
                    "You are an institutional monitoring and "
                    "evaluation officer writing formal Kenya "
                    "School of Government evaluation reports. "
                    "Write in a professional, concise, "
                    "evidence-based and human tone."
                )
            },

            {
                "role": "user",

                "content": prompt
            }

        ]
    )

    return response.choices[0].message.content.strip()


# =========================================================
# REPORT FUNCTION
# =========================================================

def generate_fe_report_llm(

    cleaned_file,

    programme_title,

    programme_code,

    duration,

    venue,

    coordinator,

    assistant,

    total_participants,

    output_folder="outputs"

):

    # =====================================================
    # CREATE OUTPUT FOLDER
    # =====================================================

    os.makedirs(
        output_folder,
        exist_ok=True
    )


    # =====================================================
    # LOAD DATA
    # =====================================================

    df = pd.read_excel(
        cleaned_file
    )

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
        .str.title()
    )


    # =====================================================
    # STANDARDIZE COLUMN NAMES
    # =====================================================

    rename_map = {

        "Topic": "Topic Description",

        "Session Topic": "Topic Description",

        "Facilitator": "Lecturer Name",

        "Lecturer": "Lecturer Name"
    }

    df = df.rename(
        columns=rename_map
    )


    # =====================================================
    # SORT DATA
    # =====================================================

    df = df.sort_values(
        by=[
            "Lecturer Name",
            "Topic Description"
        ]
    )


    # =====================================================
    # RATING COLUMNS
    # =====================================================

    rating_cols = [

        "Punctuality",

        "Presentation Flow",

        "Handling Questions",

        "Active Participation Of Learners",

        "Use Of Visual Aids",

        "Relevance Of Subject To Workplace",

        "Use Of Relevant Examples",

        "Knowledge Of Subject",

        "Treats Participants With Dignity And Respect",

        "Variety And Appropriateness Of Training Methods"
    ]

    rating_cols = [

        col

        for col in rating_cols

        if col in df.columns
    ]


    # =====================================================
    # CREATE DOCUMENT
    # =====================================================

    doc = Document()

    style = doc.styles["Normal"]

    style.font.name = "Times New Roman"

    style.font.size = Pt(11)


    # =====================================================
    # GROUP BY FACILITATOR
    # =====================================================

    grouped_facilitators = df.groupby(
        "Lecturer Name"
    )


    # =====================================================
    # LOOP THROUGH FACILITATORS
    # =====================================================

    for facilitator, facilitator_data in grouped_facilitators:


        # -------------------------------------------------
        # HEADER
        # -------------------------------------------------

        doc.add_paragraph(
            "KSG/17/FER/08"
        )

        doc.add_paragraph(
            "KENYA SCHOOL OF GOVERNMENT"
        )

        doc.add_paragraph(
            "MATUGA"
        )

        doc.add_heading(
            "FACILITATOR EVALUATION REPORT",
            level=1
        )


        # -------------------------------------------------
        # DETAILS TABLE
        # -------------------------------------------------

        details_table = doc.add_table(
            rows=4,
            cols=2
        )

        details_table.cell(0, 0).text = (
            "PROGRAM TITLE:"
        )

        details_table.cell(0, 1).text = (
            str(programme_title)
        )

        details_table.cell(1, 0).text = (
            "PROGRAMME CODE:"
        )

        details_table.cell(1, 1).text = (
            str(programme_code)
        )

        details_table.cell(2, 0).text = (
            "FACILITATOR:"
        )

        details_table.cell(2, 1).text = (
            str(facilitator)
        )

        details_table.cell(3, 0).text = (
            "DURATION:"
        )

        details_table.cell(3, 1).text = (
            str(duration)
        )

        set_table_borders(
            details_table
        )


        # -------------------------------------------------
        # INTRODUCTION
        # -------------------------------------------------

        doc.add_paragraph()

        doc.add_paragraph(
            "This report provides information to KSG "
            "management for decision making and action. "
            "The Facilitator Evaluation forms are filled "
            "by participants during the course of training "
            "programmes. The Head of Department – Training "
            "is expected to discuss the evaluation results "
            "with individual facilitators where necessary."
        )


        # -------------------------------------------------
        # PARTICIPANTS RATINGS
        # -------------------------------------------------

        doc.add_heading(
            "I. Participants' Ratings",
            level=2
        )

        doc.add_paragraph(
            "The rating distribution is presented as "
            "5/4/3/2/1, representing the number of "
            "participants selecting ratings 5, 4, 3, 2 "
            "and 1 respectively."
        )


        # =================================================
        # GET ALL SESSIONS FOR THIS FACILITATOR
        # =================================================

        sessions = list(
            facilitator_data.groupby(
                "Topic Description"
            )
        )


        # =================================================
        # CREATE CONSOLIDATED TABLE
        # =================================================

        total_columns = (
            1 +
            (len(sessions) * 2) +
            1
        )

        table = doc.add_table(
            rows=2,
            cols=total_columns
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )


        # =================================================
        # HEADER ROW 1
        # =================================================

        table.cell(
            0,
            0
        ).text = "SPECIFIC ASPECTS"

        format_cell(
            table.cell(0, 0),
            bold=True
        )

        col_position = 1


        for session, group in sessions:

            start_cell = table.cell(
                0,
                col_position
            )

            end_cell = table.cell(
                0,
                col_position + 1
            )

            merged = start_cell.merge(
                end_cell
            )

            merged.text = str(session)

            format_cell(
                merged,
                bold=True
            )

            col_position += 2


        table.cell(
            0,
            col_position
        ).text = "AGGREGATE MEAN"

        format_cell(
            table.cell(
                0,
                col_position
            ),
            bold=True
        )


        # =================================================
        # HEADER ROW 2
        # =================================================

        table.cell(
            1,
            0
        ).text = ""

        col_position = 1


        for session, group in sessions:

            table.cell(
                1,
                col_position
            ).text = "5/4/3/2/1"

            format_cell(
                table.cell(
                    1,
                    col_position
                ),
                bold=True
            )

            table.cell(
                1,
                col_position + 1
            ).text = "MEAN"

            format_cell(
                table.cell(
                    1,
                    col_position + 1
                ),
                bold=True
            )

            col_position += 2


        table.cell(
            1,
            col_position
        ).text = "MEAN"

        format_cell(
            table.cell(
                1,
                col_position
            ),
            bold=True
        )


        # =================================================
        # ANALYSE ALL SESSIONS
        # =================================================

        session_results = {}

        for session, group in sessions:

            session_results[session] = (
                get_session_statistics(
                    group,
                    rating_cols
                )
            )


        # =================================================
        # TOTAL PARTICIPANTS ROW
        # =================================================

        row = table.add_row().cells

        row[0].text = (
            "NO. OF PAX / TOTAL PARTICIPANTS"
        )

        format_cell(
            row[0],
            bold=True
        )

        col_position = 1

        for session, group in sessions:

            row[col_position].text = str(
                total_participants
            )

            format_cell(
                row[col_position],
                bold=True
            )

            row[
                col_position + 1
            ].text = ""

            col_position += 2

        row[col_position].text = ""


        # =================================================
        # TOTAL RESPONDENTS ROW
        # =================================================

        row = table.add_row().cells

        row[0].text = (
            "TOTAL RESPONDENTS"
        )

        format_cell(
            row[0],
            bold=True
        )

        col_position = 1

        for session, group in sessions:

            respondents = session_results[
                session
            ][
                "total_respondents"
            ]

            row[col_position].text = str(
                respondents
            )

            format_cell(
                row[col_position],
                bold=True
            )

            row[
                col_position + 1
            ].text = ""

            col_position += 2

        row[col_position].text = ""


        # =================================================
        # RATING ROWS
        # =================================================

        aggregate_scores = {}

        for rating in rating_cols:

            row = table.add_row().cells

            row[0].text = rating

            format_cell(
                row[0]
            )

            col_position = 1

            rating_means = []

            for session, group in sessions:

                result = session_results[
                    session
                ][
                    "ratings"
                ][
                    rating
                ]

                row[
                    col_position
                ].text = result[
                    "distribution"
                ]

                format_cell(
                    row[col_position]
                )

                row[
                    col_position + 1
                ].text = str(
                    result[
                        "mean"
                    ]
                )

                format_cell(
                    row[
                        col_position + 1
                    ]
                )

                rating_means.append(
                    result[
                        "mean"
                    ]
                )

                col_position += 2


            # ---------------------------------------------
            # AGGREGATE MEAN
            # ---------------------------------------------

            aggregate_mean = round(
                sum(rating_means) /
                len(rating_means),
                2
            )

            aggregate_scores[
                rating
            ] = aggregate_mean

            row[col_position].text = str(
                aggregate_mean
            )

            format_cell(
                row[col_position],
                bold=True
            )


        # =================================================
        # MEAN ROW
        # =================================================

        row = table.add_row().cells

        row[0].text = "MEAN"

        format_cell(
            row[0],
            bold=True
        )

        col_position = 1

        session_means = []

        for session, group in sessions:

            session_mean = session_results[
                session
            ][
                "session_mean"
            ]

            session_means.append(
                session_mean
            )

            row[col_position].text = ""

            row[
                col_position + 1
            ].text = str(
                session_mean
            )

            format_cell(
                row[
                    col_position + 1
                ],
                bold=True
            )

            col_position += 2


        overall_mean = round(
            sum(session_means) /
            len(session_means),
            2
        )

        row[col_position].text = str(
            overall_mean
        )

        format_cell(
            row[col_position],
            bold=True
        )


        # =================================================
        # COMPOSITE SCORE ROW
        # =================================================

        row = table.add_row().cells

        row[0].text = (
            "COMPOSITE SCORE (%)"
        )

        format_cell(
            row[0],
            bold=True
        )

        col_position = 1


        for session, group in sessions:

            composite = session_results[
                session
            ][
                "composite_score"
            ]

            row[col_position].text = ""

            row[
                col_position + 1
            ].text = (
                f"{composite}%"
            )

            format_cell(
                row[
                    col_position + 1
                ],
                bold=True
            )

            col_position += 2


        overall_composite = round(
            (overall_mean / 5) * 100,
            1
        )

        row[col_position].text = (
            f"{overall_composite}%"
        )

        format_cell(
            row[col_position],
            bold=True
        )


        # =================================================
        # TABLE BORDERS
        # =================================================

        set_table_borders(
            table
        )


        # =================================================
        # STRONGEST AND WEAKEST AREA
        # =================================================

        strongest = max(
            aggregate_scores,
            key=aggregate_scores.get
        )

        weakest = min(
            aggregate_scores,
            key=aggregate_scores.get
        )


        # =================================================
        # QUALITATIVE ANALYSIS
        # =================================================

        doc.add_paragraph()

        likes_raw = ""

        suggestions_raw = ""


        if "Like" in facilitator_data.columns:

            likes_raw = "; ".join(

                facilitator_data[
                    "Like"
                ]
                .dropna()
                .astype(str)
            )


        if "Suggestions" in facilitator_data.columns:

            suggestions_raw = "; ".join(

                facilitator_data[
                    "Suggestions"
                ]
                .dropna()
                .astype(str)
            )


        combined_text = f"""
Most Liked:
{likes_raw}

Suggestions:
{suggestions_raw}
"""


        qualitative = analyze_qualitative(
            combined_text
        )

        paragraphs = qualitative.split(
            "\n\n"
        )


        # -------------------------------------------------
        # MOST LIKED
        # -------------------------------------------------

        doc.add_paragraph(
            "Most liked about the facilitator"
        )

        if len(paragraphs) > 0:

            doc.add_paragraph(
                paragraphs[0]
            )

        else:

            doc.add_paragraph(
                "No comments provided."
            )


        # -------------------------------------------------
        # SUGGESTIONS
        # -------------------------------------------------

        doc.add_paragraph()

        doc.add_paragraph(
            "Suggestions on areas of improvement"
        )

        if len(paragraphs) > 1:

            doc.add_paragraph(
                paragraphs[1]
            )

        else:

            doc.add_paragraph(
                "Participants expressed minimal suggestions "
                "for improvement."
            )


        # =================================================
        # HOD COMMENTS
        # =================================================

        doc.add_paragraph()

        doc.add_paragraph(
            "II. Head of Department – Training's Comments:"
        )

        hod_comment = generate_hod_comment(
            overall_composite,
            strongest,
            weakest
        )

        doc.add_paragraph(
            hod_comment
        )


        # =================================================
        # RECOMMENDATIONS
        # =================================================

        doc.add_paragraph()

        doc.add_paragraph(
            "III. Head of Department – Training's "
            "Proposals or Recommendations:"
        )

        recommendation = generate_recommendation(
            overall_composite,
            weakest
        )

        doc.add_paragraph(
            recommendation
        )


     
        # =================================================
        # HOD SIGNATURE
        # =================================================

        doc.add_paragraph()

        doc.add_paragraph(
        "IV. Head of Department – Training:"
       )

        doc.add_paragraph()

        doc.add_paragraph(
        "Signature: ............................................................"
       )

        doc.add_paragraph()

        doc.add_paragraph(
         "Date: ...................................................................."
)


        # =================================================
        # PAGE BREAK
        # =================================================

        doc.add_page_break()


    # =====================================================
    # SAVE REPORT
    # =====================================================

    base_name = os.path.splitext(

        os.path.basename(
            cleaned_file
        )

    )[0]


    output_file = os.path.join(

        output_folder,

        f"{base_name}_FE_Report_LLM.docx"

    )


    doc.save(
        output_file
    )


    print(
        f"✅ LLM Report saved: {output_file}"
    )


    return output_file
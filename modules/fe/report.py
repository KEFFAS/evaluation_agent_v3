import os
import pandas as pd

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# =========================================================
# TABLE BORDER FUNCTION
# =========================================================
def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for border_name in [

        "top",

        "left",

        "bottom",

        "right",

        "insideH",

        "insideV"

    ]:

        border = OxmlElement(f"w:{border_name}")

        border.set(qn("w:val"), "single")

        border.set(qn("w:sz"), "8")

        border.set(qn("w:space"), "0")

        border.set(qn("w:color"), "000000")

        borders.append(border)

    tblPr.append(borders)


# =========================================================
# REPORT FUNCTION
# =========================================================
def generate_fe_report(

    cleaned_file,

    programme_title,

    programme_code,

    duration,

    venue,

    coordinator,

    assistant,

    total_participants,

    output_folder="outputs"

):
    """
    Generate Facilitator Evaluation report.
    """

    os.makedirs(output_folder, exist_ok=True)

    # =====================================================
    # LOAD CLEANED DATA
    # =====================================================
    df = pd.read_excel(cleaned_file)

    # =====================================================
    # STANDARDIZE
    # =====================================================
    df.columns = (

        df.columns

        .astype(str)

        .str.strip()

        .str.title()

    )

    rename_map = {

        "Topic": "Topic Description",

        "Session Topic": "Topic Description",

        "Facilitator": "Lecturer Name",

        "Lecturer": "Lecturer Name"

    }

    df = df.rename(columns=rename_map)

    # =====================================================
    # SORT
    # =====================================================
    df = df.sort_values(

        by=[

            "Lecturer Name",

            "Topic Description"

        ]

    )

    # =====================================================
    # RATING COLUMNS
    # =====================================================
    rating_cols = [

        "Punctuality",

        "Presentation Flow",

        "Handling Questions",

        "Active Participation Of Learners",

        "Use Of Visual Aids",

        "Relevance Of Subject To Workplace",

        "Use Of Relevant Examples",

        "Knowledge Of Subject",

        "Treats Participants With Dignity And Respect",

        "Variety And Appropriateness Of Training Methods"

    ]

    rating_cols = [

        col

        for col in rating_cols

        if col in df.columns

    ]

    # =====================================================
    # GROUP
    # =====================================================
    grouped = df.groupby(

        [

            "Lecturer Name",

            "Topic Description"

        ]

    )

    # =====================================================
    # CREATE DOCUMENT
    # =====================================================
    doc = Document()

    style = doc.styles["Normal"]

    font = style.font

    font.name = "Times New Roman"

    font.size = Pt(11)

    # =====================================================
    # CREATE REPORT FOR EACH SESSION
    # =====================================================
    for (facilitator, session), group in grouped:

        # -------------------------------------------------
        # HEADER
        # -------------------------------------------------
        doc.add_paragraph("KSG/17/FER/08")
        doc.add_paragraph("KENYA SCHOOL OF GOVERNMENT")
        doc.add_paragraph("MATUGA")
        doc.add_heading(
            "FACILITATOR EVALUATION REPORT",
            level=1
        )

        # -------------------------------------------------
        # DETAILS TABLE
        # -------------------------------------------------
        details_table = doc.add_table(
            rows=4,
            cols=2
        )

        details_table.style = "Table Grid"

        details_table.cell(0, 0).text = "PROGRAMME TITLE:"
        details_table.cell(0, 1).text = programme_title

        details_table.cell(1, 0).text = "PROGRAMME CODE:"
        details_table.cell(1, 1).text = programme_code

        details_table.cell(2, 0).text = "SESSION TOPIC:"
        details_table.cell(2, 1).text = session

        details_table.cell(3, 0).text = "FACILITATOR:"
        details_table.cell(3, 1).text = facilitator

        set_table_borders(details_table)

        doc.add_paragraph()

        # -------------------------------------------------
        # INTRODUCTION
        # -------------------------------------------------
        doc.add_paragraph(
            "This report presents participants' evaluation of the facilitator "
            "during the training programme. The findings provide useful feedback "
            "to facilitators and management for continuous improvement of "
            "training delivery."
        )

        doc.add_paragraph()

        doc.add_heading(
            "I. Participants' Ratings",
            level=2
        )

        # -------------------------------------------------
        # RATINGS TABLE
        # -------------------------------------------------
        table = doc.add_table(
            rows=1,
            cols=10
        )

        table.style = "Table Grid"

        headers = [

            "Specific Aspect",

            "Total Participants",

            "Non Response",

            "5",

            "4",

            "3",

            "2",

            "1",

            "Valid Responses",

            "% Scores 4 & 5"

        ]

        for i, header in enumerate(headers):

            table.rows[0].cells[i].text = header

        # -------------------------------------------------
        # POPULATE TABLE
        # -------------------------------------------------
        for indicator in rating_cols:

            counts = group[indicator].value_counts().to_dict()

            count5 = counts.get(5, 0)
            count4 = counts.get(4, 0)
            count3 = counts.get(3, 0)
            count2 = counts.get(2, 0)
            count1 = counts.get(1, 0)

            total_valid = (
                count5 +
                count4 +
                count3 +
                count2 +
                count1
            )

            non_response = (
                total_participants -
                total_valid
            )

            pct45 = (
                ((count5 + count4) / total_valid * 100)
                if total_valid > 0
                else 0
            )

            row = table.add_row().cells

            values = [

                indicator,

                total_participants,

                non_response,

                count5,

                count4,

                count3,

                count2,

                count1,

                total_valid,

                round(pct45, 1)

            ]

            for i, value in enumerate(values):

                row[i].text = str(value)

        set_table_borders(table)

        doc.add_paragraph()

        # =====================================================
        # QUALITATIVE FEEDBACK
        # =====================================================
        likes = "; ".join(

            group["Like"]

            .dropna()

            .astype(str)

        )

        suggestions = "; ".join(

            group["Suggestions"]

            .dropna()

            .astype(str)

        )

        doc.add_heading(

            "II. Qualitative Feedback",

            level=2

        )

        doc.add_paragraph()

        doc.add_heading(

            "Most Liked About the Facilitator",

            level=3

        )

        if likes.strip():

            doc.add_paragraph(likes)

        else:

            doc.add_paragraph(

                "No comments provided."

            )

        doc.add_paragraph()

        doc.add_heading(

            "Suggestions for Improvement",

            level=3

        )

        if suggestions.strip():

            doc.add_paragraph(suggestions)

        else:

            doc.add_paragraph(

                "No suggestions provided."

            )

        # =====================================================
        # HEAD OF DEPARTMENT COMMENTS
        # =====================================================
        doc.add_paragraph()

        doc.add_heading(

            "III. Head of Department Comments",

            level=2

        )

        doc.add_paragraph(

            "........................................................................................................"

        )

        doc.add_paragraph()

        doc.add_heading(

            "IV. Recommendations",

            level=2

        )

        doc.add_paragraph(

            "........................................................................................................"

        )

        # =====================================================
        # SIGNATURE
        # =====================================================
        doc.add_paragraph()

        signature = doc.add_table(

            rows=3,

            cols=2

        )

        signature.style = "Table Grid"

        signature.cell(0, 0).text = "Prepared by"

        signature.cell(0, 1).text = "____________________________"

        signature.cell(1, 0).text = "Signature"

        signature.cell(1, 1).text = "____________________________"

        signature.cell(2, 0).text = "Date"

        signature.cell(2, 1).text = "____________________________"

        set_table_borders(signature)

        # =====================================================
        # PAGE BREAK
        # =====================================================
        doc.add_page_break()

    # =========================================================
    # SAVE REPORT
    # =========================================================
    base_name = os.path.splitext(

        os.path.basename(cleaned_file)

    )[0]

    output_file = os.path.join(

        output_folder,

        f"{base_name}_FE_Report.docx"

    )

    doc.save(output_file)

    print(f"✅ Report saved: {output_file}")

    return output_file


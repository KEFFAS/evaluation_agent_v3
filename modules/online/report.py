import os
import pandas as pd

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from openai import OpenAI
from dotenv import load_dotenv


# =========================================================
# LOAD ENVIRONMENT
# =========================================================
load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)


# =========================================================
# TABLE BORDER FUNCTION
# =========================================================
def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for border_name in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ]:

        border = OxmlElement(
            f"w:{border_name}"
        )

        border.set(
            qn("w:val"),
            "single"
        )

        border.set(
            qn("w:sz"),
            "8"
        )

        border.set(
            qn("w:color"),
            "000000"
        )

        borders.append(border)

    tblPr.append(borders)


# =========================================================
# LLM FUNCTION
# =========================================================
def generate_text(prompt):

    try:

        response = client.chat.completions.create(

            model="gpt-5-nano-2025-08-07",

            messages=[

                {
                    "role": "system",
                    "content": (
                        "You are an institutional Monitoring and "
                        "Evaluation Officer writing formal Kenya "
                        "School of Government evaluation reports. "
                        "Write professionally, concisely and in an "
                        "evidence-based human tone. Avoid exaggerated "
                        "claims, repetition and generic AI language."
                    )
                },

                {
                    "role": "user",
                    "content": prompt
                }

            ]

        )

        return (
            response.choices[0]
            .message.content
            .strip()
        )

    except Exception as e:

        print(f"LLM Error: {e}")

        return (
            "The findings provide evidence on participant "
            "assessment of the programme and highlight areas "
            "for continued improvement."
        )


# =========================================================
# READ ANALYSIS SHEET
# =========================================================
def read_analysis_sheet(
    analysis_file,
    sheet_name
):

    try:

        return pd.read_excel(
            analysis_file,
            sheet_name=sheet_name
        )

    except Exception as e:

        print(
            f"Could not read sheet "
            f"'{sheet_name}': {e}"
        )

        return pd.DataFrame()


# =========================================================
# FIND SHEET
# =========================================================
def find_sheet_name(
    analysis_file,
    keyword
):

    excel_file = pd.ExcelFile(
        analysis_file
    )

    for sheet in excel_file.sheet_names:

        if keyword.lower() in sheet.lower():

            return sheet

    return None


# =========================================================
# FORMAT NUMBER
# =========================================================
def format_number(value):

    if pd.isna(value):

        return "0"

    try:

        value = float(value)

        if value.is_integer():

            return str(
                int(value)
            )

        return str(
            round(value, 2)
        )

    except Exception:

        return str(value)


# =========================================================
# GET TOTAL RESPONDENTS
# =========================================================
def get_total_respondents(df):

    if df.empty:

        return 0

    if "Total" in df.columns:

        try:

            return int(
                df["Total"]
                .dropna()
                .max()
            )

        except Exception:

            pass

    return 0


# =========================================================
# ADD SINGLE RATING TABLE
# =========================================================
def add_single_rating_table(
    doc,
    data,
    title,
    labels
):

    if data.empty:

        doc.add_paragraph(
            "No data was available for this section."
        )

        return

    row_data = data.iloc[0]

    total = get_total_respondents(
        data
    )

    doc.add_paragraph(
        f"Total Respondents: {total}"
    )

    table = doc.add_table(
        rows=1,
        cols=7
    )

    headers = [

        title,

        f"5 - {labels[5]}",

        f"4 - {labels[4]}",

        f"3 - {labels[3]}",

        f"2 - {labels[2]}",

        f"1 - {labels[1]}",

        "MEAN"

    ]

    for i, header in enumerate(headers):

        table.rows[0].cells[i].text = (
            str(header)
        )

    # =====================================================
    # RESULTS ROW
    # =====================================================
    row = table.add_row().cells

    row[0].text = title

    row[1].text = format_number(
        row_data.get("5", 0)
    )

    row[2].text = format_number(
        row_data.get("4", 0)
    )

    row[3].text = format_number(
        row_data.get("3", 0)
    )

    row[4].text = format_number(
        row_data.get("2", 0)
    )

    row[5].text = format_number(
        row_data.get("1", 0)
    )

    row[6].text = format_number(
        row_data.get("Mean", 0)
    )

    # =====================================================
    # COMPOSITE SCORE
    # =====================================================
    composite_row = (
        table.add_row().cells
    )

    composite_row[0].text = (
        "COMPOSITE SCORE (%)"
    )

    composite_row[6].text = (

        f"{format_number(row_data.get('Composite Score (%)', 0))}%"

    )

    set_table_borders(table)


# =========================================================
# ADD SPECIFIC ASPECTS TABLE
# =========================================================
def add_specific_aspects_table(
    doc,
    data
):

    if data.empty:

        doc.add_paragraph(
            "No data was available for this section."
        )

        return 0, 0

    total = get_total_respondents(
        data
    )

    doc.add_paragraph(
        f"Total Respondents: {total}"
    )

    table = doc.add_table(
        rows=1,
        cols=7
    )

    headers = [

        "SPECIFIC ASPECTS",

        "5 - Excellent",

        "4 - Very Good",

        "3 - Satisfactory",

        "2 - Poor",

        "1 - Very Poor",

        "MEAN"

    ]

    for i, header in enumerate(headers):

        table.rows[0].cells[i].text = (
            header
        )

    # =====================================================
    # ASPECT ROWS
    # =====================================================
    for _, item in data.iterrows():

        row = table.add_row().cells

        row[0].text = str(
            item.get(
                "Specific Aspect",
                ""
            )
        )

        row[1].text = format_number(
            item.get("5", 0)
        )

        row[2].text = format_number(
            item.get("4", 0)
        )

        row[3].text = format_number(
            item.get("3", 0)
        )

        row[4].text = format_number(
            item.get("2", 0)
        )

        row[5].text = format_number(
            item.get("1", 0)
        )

        row[6].text = format_number(
            item.get("Mean", 0)
        )

    # =====================================================
    # OVERALL MEAN
    # =====================================================
    section_mean = round(
        pd.to_numeric(
            data["Mean"],
            errors="coerce"
        ).mean(),
        2
    )

    section_composite = round(
        (section_mean / 5) * 100,
        1
    )

    mean_row = table.add_row().cells

    mean_row[0].text = (
        "OVERALL MEAN"
    )

    mean_row[6].text = str(
        section_mean
    )

    # =====================================================
    # COMPOSITE
    # =====================================================
    composite_row = (
        table.add_row().cells
    )

    composite_row[0].text = (
        "COMPOSITE SCORE (%)"
    )

    composite_row[6].text = (
        f"{section_composite}%"
    )

    set_table_borders(table)

    return (
        section_mean,
        section_composite
    )


# =========================================================
# QUALITATIVE EXTRACTION
# =========================================================
def get_qualitative_text(
    df,
    keyword
):

    matching_cols = [

        col

        for col in df.columns

        if keyword.lower()
        in str(col).lower()

    ]

    if not matching_cols:

        return ""

    col = matching_cols[0]

    responses = (

        df[col]
        .dropna()
        .astype(str)
        .str.strip()

    )

    responses = responses[
        responses != ""
    ]

    responses = list(
        dict.fromkeys(responses)
    )

    return " ".join(
        responses
    )


# =========================================================
# MAIN REPORT FUNCTION
# =========================================================
def generate_online_report(

    cleaned_file,

    analysis_file,

    programme_title,

    programme_code,

    duration,

    venue,

    coordinator,

    assistant,

    output_folder="outputs"

):

    os.makedirs(
        output_folder,
        exist_ok=True
    )

    # =====================================================
    # LOAD CLEANED DATA
    # =====================================================
    cleaned_df = pd.read_excel(
        cleaned_file
    )

    # =====================================================
    # FIND ANALYSIS SHEETS
    # =====================================================
    overall_sheet = find_sheet_name(
        analysis_file,
        "Overall Summary"
    )

    objectives_sheet = find_sheet_name(
        analysis_file,
        "Course Objectives"
    )

    expectations_sheet = find_sheet_name(
        analysis_file,
        "Fulfilment"
    )

    aspects_sheet = find_sheet_name(
        analysis_file,
        "Specific Aspects"
    )

    future_sheet = find_sheet_name(
        analysis_file,
        "Future Attendance"
    )

    recommend_sheet = find_sheet_name(
        analysis_file,
        "Recommend KSG"
    )

    # =====================================================
    # READ SHEETS
    # =====================================================
    overall_df = (

        read_analysis_sheet(
            analysis_file,
            overall_sheet
        )

        if overall_sheet

        else pd.DataFrame()

    )

    objectives_df = (

        read_analysis_sheet(
            analysis_file,
            objectives_sheet
        )

        if objectives_sheet

        else pd.DataFrame()

    )

    expectations_df = (

        read_analysis_sheet(
            analysis_file,
            expectations_sheet
        )

        if expectations_sheet

        else pd.DataFrame()

    )

    aspects_df = (

        read_analysis_sheet(
            analysis_file,
            aspects_sheet
        )

        if aspects_sheet

        else pd.DataFrame()

    )

    future_df = (

        read_analysis_sheet(
            analysis_file,
            future_sheet
        )

        if future_sheet

        else pd.DataFrame()

    )

    recommend_df = (

        read_analysis_sheet(
            analysis_file,
            recommend_sheet
        )

        if recommend_sheet

        else pd.DataFrame()

    )

    # =====================================================
    # OVERALL SCORE
    # =====================================================
    overall_mean = 0
    overall_composite = 0

    if not overall_df.empty:

        overall_row = overall_df[

            overall_df[
                "Evaluation Area"
            ].astype(str)
            .str.contains(
                "OVERALL",
                case=False,
                na=False
            )

        ]

        if not overall_row.empty:

            overall_mean = float(

                overall_row.iloc[0][
                    "Mean Score"
                ]

            )

            overall_composite = float(

                overall_row.iloc[0][
                    "Composite Score (%)"
                ]

            )

    # =====================================================
    # CREATE DOCUMENT
    # =====================================================
    doc = Document()

    style = doc.styles["Normal"]

    style.font.name = (
        "Times New Roman"
    )

    style.font.size = Pt(11)

    # =====================================================
    # HEADER
    # =====================================================
    doc.add_paragraph(
        "KSG/17/EOEEF/07"
    )

    p = doc.add_paragraph()

    p.alignment = (
        WD_PARAGRAPH_ALIGNMENT.CENTER
    )

    run = p.add_run(

        "KENYA SCHOOL OF GOVERNMENT\n"
        "MATUGA\n\n"
        "ONLINE END-OF-EVENT EVALUATION REPORT"

    )

    run.bold = True

    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(16)

    # =====================================================
    # PROGRAMME DETAILS
    # =====================================================
    details_table = doc.add_table(
        rows=3,
        cols=4
    )

    details = [

        (
            0,
            0,
            "PROGRAMME TITLE:"
        ),

        (
            0,
            1,
            programme_title
        ),

        (
            0,
            2,
            "DURATION:"
        ),

        (
            0,
            3,
            duration
        ),

        (
            1,
            0,
            "PROGRAMME CODE:"
        ),

        (
            1,
            1,
            programme_code
        ),

        (
            1,
            2,
            "VENUE:"
        ),

        (
            1,
            3,
            venue
        ),

        (
            2,
            0,
            "COORDINATOR:"
        ),

        (
            2,
            1,
            coordinator
        ),

        (
            2,
            2,
            "PROGRAM ASSISTANT:"
        ),

        (
            2,
            3,
            assistant
        )

    ]

    for row, col, value in details:

        details_table.cell(
            row,
            col
        ).text = str(value)

    set_table_borders(
        details_table
    )

    # =====================================================
    # INTRODUCTION
    # =====================================================
    doc.add_heading(
        "A. PROGRAMME EVALUATION",
        level=2
    )

    doc.add_paragraph(

        "The Kenya School of Government conducted an "
        "online end-of-event evaluation to assess participant "
        "perceptions of the quality, relevance and effectiveness "
        "of the training programme. The evaluation focused on "
        "achievement of course objectives, fulfilment of "
        "expectations, specific aspects of programme delivery, "
        "future participation and willingness to recommend "
        "KSG online learning."

    )

    # =====================================================
    # SECTION 1
    # OVERALL SUMMARY
    # =====================================================
    doc.add_heading(
        "1. Overall Evaluation Summary",
        level=2
    )

    if not overall_df.empty:

        summary_table = doc.add_table(
            rows=1,
            cols=3
        )

        headers = [

            "EVALUATION AREA",

            "MEAN SCORE",

            "COMPOSITE SCORE (%)"

        ]

        for i, header in enumerate(headers):

            summary_table.rows[0].cells[i].text = (
                header
            )

        for _, item in overall_df.iterrows():

            row = summary_table.add_row().cells

            row[0].text = str(
                item.get(
                    "Evaluation Area",
                    ""
                )
            )

            row[1].text = format_number(
                item.get(
                    "Mean Score",
                    0
                )
            )

            row[2].text = (

                f"{format_number(item.get('Composite Score (%)', 0))}%"

            )

        set_table_borders(
            summary_table
        )

    summary_prompt = f"""
Interpret the following overall results from a
Kenya School of Government online training
programme.

Overall Mean Score:
{overall_mean}

Overall Composite Score:
{overall_composite}%

Write ONE concise institutional paragraph.

Mention the overall performance level and major
strengths. Mention areas requiring attention only
where supported by the evaluation results.
Do not exaggerate.
"""

    doc.add_paragraph(
        generate_text(
            summary_prompt
        )
    )

    # =====================================================
    # SECTION 2
    # COURSE OBJECTIVES
    # =====================================================
    doc.add_heading(
        "2. Course Objectives Achievement",
        level=2
    )

    objective_labels = {

        5: "Excellent",

        4: "Very Good",

        3: "Satisfactory",

        2: "Poor",

        1: "Very Poor"

    }

    add_single_rating_table(

        doc,

        objectives_df,

        "COURSE OBJECTIVES ACHIEVEMENT",

        objective_labels

    )

    if not objectives_df.empty:

        item = objectives_df.iloc[0]

        prompt = f"""
Interpret these results on Course Objectives Achievement.

Mean Score:
{item.get('Mean', 0)}

Composite Score:
{item.get('Composite Score (%)', 0)}%

Write ONE concise professional paragraph.
"""

        doc.add_paragraph(
            generate_text(prompt)
        )

    # =====================================================
    # SECTION 3
    # PERSONAL EXPECTATIONS
    # =====================================================
    doc.add_heading(
        "3. Fulfilment of Personal Expectations",
        level=2
    )

    expectation_labels = {

        5: "Great Extent",

        4: "Some Extent",

        3: "Satisfactory",

        2: "Not Sure",

        1: "Not At All"

    }

    add_single_rating_table(

        doc,

        expectations_df,

        "FULFILMENT OF PERSONAL EXPECTATIONS",

        expectation_labels

    )

    if not expectations_df.empty:

        item = expectations_df.iloc[0]

        prompt = f"""
Interpret these results on fulfilment of
participant expectations.

Mean Score:
{item.get('Mean', 0)}

Composite Score:
{item.get('Composite Score (%)', 0)}%

Write ONE concise institutional paragraph.
"""

        doc.add_paragraph(
            generate_text(prompt)
        )

    # =====================================================
    # SECTION 4
    # SPECIFIC ASPECTS
    # =====================================================
    doc.add_heading(

        "4. Ratings on Specific Aspects of the Online Training Programme",

        level=2

    )

    aspects_mean, aspects_composite = (
        add_specific_aspects_table(
            doc,
            aspects_df
        )
    )

    if not aspects_df.empty:

        findings = []

        for _, item in aspects_df.iterrows():

            findings.append(

                f"{item.get('Specific Aspect', '')}: "
                f"Mean {item.get('Mean', 0)}, "
                f"Composite "
                f"{item.get('Composite Score (%)', 0)}%"

            )

        prompt = f"""
Interpret the following specific aspect ratings
from a Kenya School of Government online training
programme:

{chr(10).join(findings)}

Overall Mean:
{aspects_mean}

Overall Composite Score:
{aspects_composite}%

Write ONE concise institutional paragraph.

Identify major strengths and comparatively
lower-rated areas strictly based on the data.
Do not exaggerate or repeat every score.
"""

        doc.add_paragraph(
            generate_text(prompt)
        )

    # =====================================================
    # SECTION 5
    # FUTURE ATTENDANCE
    # =====================================================
    doc.add_heading(

        "5. Likelihood of Attending Future Online Training Sessions at KSG",

        level=2

    )

    future_labels = {

        5: "Great Extent",

        4: "Some Extent",

        3: "Satisfactory",

        2: "Not Sure",

        1: "Not At All"

    }

    add_single_rating_table(

        doc,

        future_df,

        "LIKELIHOOD OF FUTURE ATTENDANCE",

        future_labels

    )

    if not future_df.empty:

        item = future_df.iloc[0]

        prompt = f"""
Interpret these results regarding willingness to
attend future KSG online training programmes.

Mean Score:
{item.get('Mean', 0)}

Composite Score:
{item.get('Composite Score (%)', 0)}%

Write ONE concise professional paragraph.
"""

        doc.add_paragraph(
            generate_text(prompt)
        )

    # =====================================================
    # SECTION 6
    # RECOMMEND KSG
    # =====================================================
    doc.add_heading(

        "6. Willingness to Recommend KSG Online Learning to Colleagues and Friends",

        level=2

    )

    recommend_labels = {

        5: "Great Extent",

        4: "Some Extent",

        3: "Satisfactory",

        2: "Not Sure",

        1: "Not At All"

    }

    add_single_rating_table(

        doc,

        recommend_df,

        "WILLINGNESS TO RECOMMEND KSG ONLINE LEARNING",

        recommend_labels

    )

    if not recommend_df.empty:

        item = recommend_df.iloc[0]

        prompt = f"""
Interpret these results regarding participants'
willingness to recommend KSG online learning.

Mean Score:
{item.get('Mean', 0)}

Composite Score:
{item.get('Composite Score (%)', 0)}%

Write ONE concise professional paragraph.
"""

        doc.add_paragraph(
            generate_text(prompt)
        )

    # =====================================================
    # QUALITATIVE SECTIONS
    # =====================================================
    qualitative_sections = [

        (
            "7. Suggestions for Improving KSG eLearning Courses and Enhancing Participants' Learning Experiences",
            "improv"
        ),

        (
            "8. Additional Comments and Suggestions on the Training Programme",
            "comment"
        ),

        (
            "9. Recommended Additional Topics for Inclusion in the Training Programme",
            "topic"
        ),

        (
            "10. Additional Training Programmes of Interest",
            "interest"
        )

    ]

    qualitative_data = {}

    for title, keyword in qualitative_sections:

        raw_text = get_qualitative_text(
            cleaned_df,
            keyword
        )

        qualitative_data[keyword] = raw_text

        doc.add_heading(
            title,
            level=2
        )

        if raw_text.strip():

            prompt = f"""
The following are participant responses from an
online Kenya School of Government training
evaluation:

{raw_text}

Write ONE concise institutional paragraph that:

- Summarizes recurring themes.
- Uses professional and natural language.
- Avoids bullets.
- Avoids repetition.
- Reflects only what participants said.
- Does not invent information.
"""

            doc.add_paragraph(
                generate_text(prompt)
            )

        else:

            doc.add_paragraph(
                "No responses were provided for this section."
            )

    # =====================================================
    # SECTION 11
    # KEY INSIGHTS AND RECOMMENDATIONS
    # =====================================================
    doc.add_heading(
        "11. Key Insights and Recommendations",
        level=2
    )

    quantitative_findings = ""

    if not overall_df.empty:

        for _, item in overall_df.iterrows():

            quantitative_findings += (

                f"{item.get('Evaluation Area', '')}: "

                f"Mean={item.get('Mean Score', 0)}, "

                f"Composite="
                f"{item.get('Composite Score (%)', 0)}%\n"

            )

    qualitative_findings = "\n\n".join(

        [

            f"{key}: {value}"

            for key, value
            in qualitative_data.items()

            if value.strip()

        ]

    )

    prompt = f"""
Using ONLY the following findings, prepare Key
Insights and Recommendations for a Kenya School
of Government online training evaluation report.

QUANTITATIVE FINDINGS:

{quantitative_findings}

QUALITATIVE FEEDBACK:

{qualitative_findings}

Requirements:

1. Provide 3 to 5 Key Insights.
2. Provide 3 to 5 Recommendations.
3. Base everything strictly on the findings.
4. Do not invent concerns.
5. Keep recommendations practical.
6. Use concise institutional language.

Format exactly as:

Key Insights
• ...
• ...
• ...

Recommendations
• ...
• ...
• ...
"""

    doc.add_paragraph(
        generate_text(prompt)
    )

    # =====================================================
    # SIGNATURES
    # =====================================================
    doc.add_paragraph()

    doc.add_paragraph(
        "Prepared by............................................................"
    )

    doc.add_paragraph(
        "Date............................   "
        "Signature............................"
    )

    doc.add_paragraph()

    doc.add_paragraph(
        "Confirmed by..........................................................."
    )

    doc.add_paragraph(
        "Date............................   "
        "Signature............................"
    )

    doc.add_paragraph()

    doc.add_paragraph(
        "Approved by............................................................"
    )

    doc.add_paragraph(
        "Date............................   "
        "Signature............................"
    )

    # =====================================================
    # SAVE REPORT
    # =====================================================
    base_name = os.path.splitext(
        os.path.basename(cleaned_file)
    )[0]

    output_file = os.path.join(

        output_folder,

        f"{base_name}_Online_EEE_Report.docx"

    )

    doc.save(
        output_file
    )

    print("=" * 60)
    print("ONLINE EEE REPORT GENERATED")
    print("=" * 60)
    print(
        f"Report saved to: {output_file}"
    )

    return output_file
import os
import re
import pandas as pd
import numpy as np

from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =========================================================
# CONSTANTS
# =========================================================

RATING_COLS = [
    "Punctuality",
    "Presentation Flow",
    "Handling Questions",
    "Active Participation Of Learners",
    "Use Of Visual Aids",
    "Relevance Of Subject To Workplace",
    "Use Of Relevant Examples",
    "Knowledge Of Subject",
    "Treats Participants With Dignity And Respect",
    "Variety And Appropriateness Of Training Methods"
]


# =========================================================
# CLEAN SHEET NAME
# =========================================================

def clean_sheet_name(name):
    """
    Make a valid Excel worksheet name.
    """

    name = str(name)

    # Replace invalid Excel characters
    name = re.sub(
        r'[:\\/*?\[\]]',
        '-',
        name
    )

    name = name.strip()

    # Excel sheet name limit
    return name[:31]


# =========================================================
# CLEAN RATING SERIES
# =========================================================

def clean_rating_series(series):
    """
    Converts ratings to numeric values and keeps
    only valid ratings between 1 and 5.
    """

    ratings = pd.to_numeric(
        series,
        errors="coerce"
    )

    return ratings.where(
        ratings.isin([1, 2, 3, 4, 5])
    )


# =========================================================
# CALCULATE TOTAL RESPONDENTS
# =========================================================

def calculate_total_respondents(
    group,
    rating_cols
):
    """
    Counts participants who responded to at least
    one rating question.
    """

    available_cols = [

        col

        for col in rating_cols

        if col in group.columns
    ]

    if not available_cols:

        return 0

    ratings = group[
        available_cols
    ].copy()

    for col in available_cols:

        ratings[col] = (
            clean_rating_series(
                ratings[col]
            )
        )

    respondents = ratings.notna().any(
        axis=1
    )

    return int(
        respondents.sum()
    )


# =========================================================
# CALCULATE RATING STATISTICS
# =========================================================

def calculate_rating_stats(
    group,
    column
):
    """
    Calculates:

    - Rating distribution 5/4/3/2/1
    - Mean score
    - Total valid responses
    """

    ratings = clean_rating_series(
        group[column]
    )

    count_5 = int(
        (ratings == 5).sum()
    )

    count_4 = int(
        (ratings == 4).sum()
    )

    count_3 = int(
        (ratings == 3).sum()
    )

    count_2 = int(
        (ratings == 2).sum()
    )

    count_1 = int(
        (ratings == 1).sum()
    )

    total_valid = (
        count_5
        + count_4
        + count_3
        + count_2
        + count_1
    )

    # =====================================================
    # CALCULATE MEAN
    # =====================================================

    if total_valid > 0:

        mean_score = (

            (5 * count_5)
            + (4 * count_4)
            + (3 * count_3)
            + (2 * count_2)
            + (1 * count_1)

        ) / total_valid

    else:

        mean_score = 0

    distribution = (

        f"{count_5}/"
        f"{count_4}/"
        f"{count_3}/"
        f"{count_2}/"
        f"{count_1}"

    )

    return {

        "distribution":
            distribution,

        "mean":
            round(mean_score, 2),

        "total_valid":
            total_valid
    }


# =========================================================
# ANALYSE ONE SESSION
# =========================================================

def analyse_session(
    group,
    rating_cols,
    total_participants
):
    """
    Analyses one facilitator session.
    """

    # =====================================================
    # TOTAL RESPONDENTS
    # =====================================================

    total_respondents = (
        calculate_total_respondents(
            group,
            rating_cols
        )
    )

    # =====================================================
    # INDIVIDUAL RATINGS
    # =====================================================

    ratings = {}

    means = []

    for col in rating_cols:

        if col not in group.columns:

            continue

        stats = calculate_rating_stats(
            group,
            col
        )

        ratings[col] = stats

        if stats["total_valid"] > 0:

            means.append(
                stats["mean"]
            )

    # =====================================================
    # SESSION MEAN
    # =====================================================

    if means:

        session_mean = round(
            np.mean(means),
            2
        )

    else:

        session_mean = 0

    # =====================================================
    # COMPOSITE SCORE
    # =====================================================

    composite_score = round(

        (session_mean / 5) * 100,

        1

    ) if session_mean > 0 else 0

    return {

        "total_participants":
            total_participants,

        "total_respondents":
            total_respondents,

        "ratings":
            ratings,

        "session_mean":
            session_mean,

        "composite_score":
            composite_score
    }


# =========================================================
# ANALYSE FACILITATOR
# =========================================================

def analyse_facilitator(
    facilitator_data,
    rating_cols,
    total_participants
):
    """
    Analyses all sessions belonging to one facilitator.
    """

    sessions = {}

    # =====================================================
    # GROUP BY SESSION
    # =====================================================

    grouped_sessions = (

        facilitator_data.groupby(
            "Topic Description",
            sort=False
        )

    )

    # =====================================================
    # ANALYSE EACH SESSION
    # =====================================================

    for session, group in grouped_sessions:

        sessions[str(session)] = (

            analyse_session(

                group,

                rating_cols,

                total_participants
            )

        )

    # =====================================================
    # AGGREGATE MEAN FOR EACH ASPECT
    # =====================================================

    aggregate_means = {}

    for rating in rating_cols:

        rating_means = []

        for session_data in sessions.values():

            stats = (

                session_data[
                    "ratings"
                ].get(
                    rating
                )

            )

            if stats:

                if stats[
                    "total_valid"
                ] > 0:

                    rating_means.append(
                        stats[
                            "mean"
                        ]
                    )

        if rating_means:

            aggregate_means[
                rating
            ] = round(

                np.mean(
                    rating_means
                ),

                2
            )

        else:

            aggregate_means[
                rating
            ] = 0

    # =====================================================
    # OVERALL MEAN
    # =====================================================

    valid_session_means = [

        session[
            "session_mean"
        ]

        for session
        in sessions.values()

        if session[
            "session_mean"
        ] > 0
    ]

    if valid_session_means:

        overall_mean = round(

            np.mean(
                valid_session_means
            ),

            2
        )

    else:

        overall_mean = 0

    # =====================================================
    # OVERALL COMPOSITE SCORE
    # =====================================================

    overall_composite = round(

        (overall_mean / 5) * 100,

        1

    ) if overall_mean > 0 else 0

    # =====================================================
    # STRONGEST AND WEAKEST
    # =====================================================

    valid_scores = {

        key: value

        for key, value
        in aggregate_means.items()

        if value > 0
    }

    if valid_scores:

        strongest = max(
            valid_scores,
            key=valid_scores.get
        )

        weakest = min(
            valid_scores,
            key=valid_scores.get
        )

    else:

        strongest = None
        weakest = None

    return {

        "sessions":
            sessions,

        "aggregate_means":
            aggregate_means,

        "overall_mean":
            overall_mean,

        "overall_composite":
            overall_composite,

        "strongest":
            strongest,

        "weakest":
            weakest
    }


# =========================================================
# CREATE FACILITATOR ANALYSIS TABLE
# =========================================================

def create_facilitator_table(
    analysis,
    rating_cols
):
    """
    Creates a consolidated table for one facilitator.

    Structure:

    Specific Aspect

    Session 1:
        5/4/3/2/1
        Mean

    Session 2:
        5/4/3/2/1
        Mean

    Aggregate Mean
    """

    sessions = analysis["sessions"]

    session_names = list(
        sessions.keys()
    )

    rows = []

    # =====================================================
    # TOTAL PARTICIPANTS
    # =====================================================

    participants_row = {

        "Specific Aspects":
            "NO. OF PAX / TOTAL PARTICIPANTS"
    }

    for session_name in session_names:

        participants_row[
            f"{session_name} - 5/4/3/2/1"
        ] = sessions[
            session_name
        ][
            "total_participants"
        ]

        participants_row[
            f"{session_name} - Mean"
        ] = ""

    participants_row[
        "Aggregate Mean"
    ] = ""

    rows.append(
        participants_row
    )

    # =====================================================
    # TOTAL RESPONDENTS
    # =====================================================

    respondents_row = {

        "Specific Aspects":
            "TOTAL RESPONDENTS"
    }

    for session_name in session_names:

        respondents_row[
            f"{session_name} - 5/4/3/2/1"
        ] = sessions[
            session_name
        ][
            "total_respondents"
        ]

        respondents_row[
            f"{session_name} - Mean"
        ] = ""

    respondents_row[
        "Aggregate Mean"
    ] = ""

    rows.append(
        respondents_row
    )

    # =====================================================
    # RATING ROWS
    # =====================================================

    for rating in rating_cols:

        row = {

            "Specific Aspects":
                rating
        }

        for session_name in session_names:

            rating_data = (

                sessions[
                    session_name
                ][
                    "ratings"
                ].get(
                    rating,
                    {}
                )

            )

            row[
                f"{session_name} - 5/4/3/2/1"
            ] = rating_data.get(
                "distribution",
                "-"
            )

            row[
                f"{session_name} - Mean"
            ] = rating_data.get(
                "mean",
                "-"
            )

        row[
            "Aggregate Mean"
        ] = analysis[
            "aggregate_means"
        ].get(
            rating,
            "-"
        )

        rows.append(
            row
        )

    # =====================================================
    # MEAN ROW
    # =====================================================

    mean_row = {

        "Specific Aspects":
            "MEAN"
    }

    for session_name in session_names:

        mean_row[
            f"{session_name} - 5/4/3/2/1"
        ] = ""

        mean_row[
            f"{session_name} - Mean"
        ] = sessions[
            session_name
        ][
            "session_mean"
        ]

    mean_row[
        "Aggregate Mean"
    ] = analysis[
        "overall_mean"
    ]

    rows.append(
        mean_row
    )

    # =====================================================
    # COMPOSITE SCORE ROW
    # =====================================================

    composite_row = {

        "Specific Aspects":
            "COMPOSITE SCORE (%)"
    }

    for session_name in session_names:

        composite_row[
            f"{session_name} - 5/4/3/2/1"
        ] = ""

        composite_row[
            f"{session_name} - Mean"
        ] = (

            f"{sessions[session_name]['composite_score']}%"
        )

    composite_row[
        "Aggregate Mean"
    ] = (

        f"{analysis['overall_composite']}%"
    )

    rows.append(
        composite_row
    )

    return pd.DataFrame(
        rows
    )


# =========================================================
# AUTO WIDTH
# =========================================================

def auto_adjust_width(
    worksheet
):
    """
    Adjusts column widths.
    """

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = (
            get_column_letter(
                column_cells[0].column
            )
        )

        for cell in column_cells:

            try:

                cell_length = len(
                    str(cell.value)
                )

                if cell_length > max_length:

                    max_length = cell_length

            except Exception:

                pass

        adjusted_width = min(
            max_length + 2,
            40
        )

        worksheet.column_dimensions[
            column_letter
        ].width = adjusted_width


# =========================================================
# FORMAT WORKSHEET
# =========================================================

def format_worksheet(
    worksheet
):
    """
    Applies professional formatting.
    """

    thin = Side(
        style="thin"
    )

    border = Border(

        left=thin,

        right=thin,

        top=thin,

        bottom=thin
    )

    # =====================================================
    # APPLY FONT AND ALIGNMENT
    # =====================================================

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(

                horizontal="center",

                vertical="center",

                wrap_text=True
            )

            cell.border = border

    # =====================================================
    # FIRST COLUMN LEFT ALIGNMENT
    # =====================================================

    for row in worksheet.iter_rows():

        row[0].alignment = Alignment(

            horizontal="left",

            vertical="center",

            wrap_text=True
        )

    # =====================================================
    # BOLD TITLE ROWS
    # =====================================================

    for row in range(
        1,
        worksheet.max_row + 1
    ):

        value = worksheet.cell(
            row=row,
            column=1
        ).value

        if value in [

            "NO. OF PAX / TOTAL PARTICIPANTS",

            "TOTAL RESPONDENTS",

            "MEAN",

            "COMPOSITE SCORE (%)"
        ]:

            for cell in worksheet[row]:

                cell.font = Font(
                    bold=True
                )

    auto_adjust_width(
        worksheet
    )


# =========================================================
# CREATE SUMMARY SHEET
# =========================================================

def create_summary_sheet(
    writer,
    summary_rows
):
    """
    Creates Facilitator Summary sheet.
    """

    summary_df = pd.DataFrame(

        summary_rows,

        columns=[

            "Facilitator",

            "Number of Sessions",

            "Overall Mean",

            "Composite Score (%)",

            "Strongest Aspect",

            "Weakest Aspect"
        ]
    )

    summary_df.to_excel(

        writer,

        sheet_name="Facilitator Summary",

        index=False
    )

    worksheet = writer.sheets[
        "Facilitator Summary"
    ]

    # =====================================================
    # FORMAT HEADER
    # =====================================================

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(

            horizontal="center",

            vertical="center",

            wrap_text=True
        )

    # =====================================================
    # FORMAT CELLS
    # =====================================================

    thin = Side(
        style="thin"
    )

    border = Border(

        left=thin,

        right=thin,

        top=thin,

        bottom=thin
    )

    for row in worksheet.iter_rows():

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(

                horizontal="center",

                vertical="center",

                wrap_text=True
            )

    auto_adjust_width(
        worksheet
    )


# =========================================================
# MAIN ANALYSIS FUNCTION
# =========================================================

def analyze_fe(

    cleaned_file,

    total_participants,

    output_folder="outputs"
):
    """
    Analyse Facilitator Evaluation data.

    Produces:

    1. Facilitator Summary Sheet
    2. One consolidated sheet per Facilitator

    Returns
    -------

    str
        Path to analysed workbook.
    """

    # =====================================================
    # CREATE OUTPUT FOLDER
    # =====================================================

    os.makedirs(

        output_folder,

        exist_ok=True
    )

    # =====================================================
    # LOAD DATA
    # =====================================================

    df = pd.read_excel(
        cleaned_file
    )

    # =====================================================
    # STANDARDIZE COLUMN NAMES
    # =====================================================

    df.columns = (

        df.columns

        .astype(str)

        .str.strip()

        .str.title()
    )

    rename_map = {

        "Topic":
            "Topic Description",

        "Session Topic":
            "Topic Description",

        "Facilitator":
            "Lecturer Name",

        "Lecturer":
            "Lecturer Name"
    }

    df = df.rename(
        columns=rename_map
    )

    # =====================================================
    # VALIDATION
    # =====================================================

    required = [

        "Topic Description",

        "Lecturer Name"
    ]

    missing = [

        col

        for col in required

        if col not in df.columns
    ]

    if missing:

        raise ValueError(

            f"Missing required columns: {missing}"
        )

    # =====================================================
    # AVAILABLE RATING COLUMNS
    # =====================================================

    rating_cols = [

        col

        for col in RATING_COLS

        if col in df.columns
    ]

    if not rating_cols:

        raise ValueError(
            "No rating columns found."
        )

    # =====================================================
    # SORT DATA
    # =====================================================

    df = df.sort_values(

        by=[

            "Lecturer Name",

            "Topic Description"
        ]

    )

    # =====================================================
    # OUTPUT FILE
    # =====================================================

    base_name = os.path.splitext(

        os.path.basename(
            cleaned_file
        )

    )[0]

    output_file = os.path.join(

        output_folder,

        f"{base_name}_analysis.xlsx"
    )

    # =====================================================
    # SUMMARY DATA
    # =====================================================

    summary_rows = []

    # =====================================================
    # WRITE EXCEL FILE
    # =====================================================

    with pd.ExcelWriter(

        output_file,

        engine="openpyxl"
    ) as writer:

        # =================================================
        # GROUP BY FACILITATOR
        # =================================================

        grouped = df.groupby(

            "Lecturer Name",

            sort=False
        )

        for facilitator, facilitator_data in grouped:

            # =============================================
            # ANALYSE FACILITATOR
            # =============================================

            analysis = analyse_facilitator(

                facilitator_data,

                rating_cols,

                total_participants
            )

            # =============================================
            # CREATE ANALYSIS TABLE
            # =============================================

            result_df = create_facilitator_table(

                analysis,

                rating_cols
            )

            # =============================================
            # SHEET NAME
            # =============================================

            sheet_name = clean_sheet_name(
                facilitator
            )

            # Avoid duplicate sheet names

            if sheet_name in writer.book.sheetnames:

                sheet_name = clean_sheet_name(
                    f"{facilitator[:25]}_2"
                )

            # =============================================
            # WRITE FACILITATOR DETAILS
            # =============================================

            details = pd.DataFrame({

                "Facilitator Evaluation Analysis": [

                    f"FACILITATOR: {facilitator}",

                    f"TOTAL PARTICIPANTS: {total_participants}",

                    f"NUMBER OF SESSIONS: {len(analysis['sessions'])}",

                    "RATING DISTRIBUTION: 5/4/3/2/1"
                ]

            })

            details.to_excel(

                writer,

                sheet_name=sheet_name,

                index=False,

                header=False,

                startrow=0
            )

            # =============================================
            # WRITE ANALYSIS TABLE
            # =============================================

            result_df.to_excel(

                writer,

                sheet_name=sheet_name,

                index=False,

                startrow=5
            )

            worksheet = writer.sheets[
                sheet_name
            ]

            # =============================================
            # FORMAT DETAILS
            # =============================================

            for row in range(1, 5):

                worksheet.cell(
                    row=row,
                    column=1
                ).font = Font(
                    bold=True
                )

            # =============================================
            # FORMAT TABLE HEADER
            # =============================================

            header_row = 6

            for cell in worksheet[
                header_row
            ]:

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(

                    horizontal="center",

                    vertical="center",

                    wrap_text=True
                )

            # =============================================
            # FORMAT SHEET
            # =============================================

            format_worksheet(
                worksheet
            )

            # =============================================
            # ADD SUMMARY
            # =============================================

            summary_rows.append({

                "Facilitator":
                    facilitator,

                "Number of Sessions":
                    len(
                        analysis[
                            "sessions"
                        ]
                    ),

                "Overall Mean":
                    analysis[
                        "overall_mean"
                    ],

                "Composite Score (%)":
                    analysis[
                        "overall_composite"
                    ],

                "Strongest Aspect":
                    analysis[
                        "strongest"
                    ],

                "Weakest Aspect":
                    analysis[
                        "weakest"
                    ]
            })

        # =================================================
        # CREATE SUMMARY SHEET
        # =================================================

        create_summary_sheet(

            writer,

            summary_rows
        )

    print(
        f"Analysis saved: {output_file}"
    )

    return output_file
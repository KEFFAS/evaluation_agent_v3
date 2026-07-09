# =========================================================
# ONLINE END-OF-EVENT REPORT GENERATOR (LLM VERSION)
# FINAL INSTITUTIONAL VERSION
# =========================================================

import pandas as pd
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openai import OpenAI
from dotenv import load_dotenv

# =========================================================
# LOAD ENV VARIABLES
# =========================================================
load_dotenv()

# =========================================================
# OPENAI CLIENT
# =========================================================
client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)

# =========================================================
# TABLE BORDER FUNCTION
# =========================================================
def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement('w:tblBorders')

    for border_name in [
        'top',
        'left',
        'bottom',
        'right',
        'insideH',
        'insideV'
    ]:

        border = OxmlElement(f'w:{border_name}')

        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:color'), '000000')

        borders.append(border)

    tblPr.append(borders)

# =========================================================
# LLM FUNCTION
# =========================================================
def generate_text(prompt):

    response = client.chat.completions.create(

        model="gpt-5-nano-2025-08-07",

        messages=[
            {
                "role":"system",
                "content":
                "You are an institutional monitoring and evaluation officer "
                "writing formal Kenya School of Government evaluation reports. "
                "Write in a professional, concise, evidence-based and human tone. "
                "Avoid exaggerated language, repetition, and generic AI wording."
            },
            {
                "role":"user",
                "content":prompt
            }
        ]
    )

    return response.choices[0].message.content.strip()

# =========================================================
# TABLE FUNCTION
# =========================================================
def add_percentage_table(doc, column_name, labels):

    table = doc.add_table(rows=1, cols=2)

    table.rows[0].cells[0].text = "Rating"
    table.rows[0].cells[1].text = "Percentage of Respondents"

    counts = df[column_name].value_counts().to_dict()

    total = sum(counts.values())

    percentages = {}

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        percentages[score] = pct

        row = table.add_row().cells

        row[0].text = labels[score]
        row[1].text = str(pct)

    set_table_borders(table)

    return percentages

# =========================================================
# QUALITATIVE EXTRACTION
# =========================================================
def get_qualitative_text(keyword):

    matching_cols = [

        col for col in columns

        if keyword in col.lower()
    ]

    if matching_cols:

        col = matching_cols[0]

        responses = (
            df[col]
            .dropna()
            .astype(str)
            .str.strip()
        )

        responses = responses[
            responses != ""
        ]

        responses = list(dict.fromkeys(responses))

        return " ".join(responses)

    return ""

# =========================================================
# INPUT FILE
# =========================================================
file_name = input(
    "Enter cleaned Online EEE file: "
).strip()

# =========================================================
# LOAD DATA
# =========================================================
df = pd.read_excel(
    file_name,
    engine="openpyxl"
)

print("\nLoaded cleaned dataset successfully.")

# =========================================================
# USER INPUTS
# =========================================================
program_title = input("Enter Program Title: ")
duration = input("Enter Program Duration: ")
program_code = input("Enter Program Code: ")
venue = input("Enter Venue: ")
coordinator = input("Enter Coordinator Name: ")
assistant = input("Enter Program Assistant Name: ")

# =========================================================
# CREATE DOCUMENT
# =========================================================
doc = Document()

style = doc.styles['Normal']

style.font.name = 'Times New Roman'
style.font.size = Pt(11)

# =========================================================
# HEADER
# =========================================================
doc.add_paragraph("KSG/17/EOEEF/07")

p = doc.add_paragraph()

p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

run = p.add_run(
    "KENYA SCHOOL OF GOVERNMENT\n"
    "MATUGA\n\n"
    "END-OF-EVENT EVALUATION REPORT"
)

run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(16)

# =========================================================
# PROGRAM DETAILS TABLE
# =========================================================
details_table = doc.add_table(
    rows=3,
    cols=4
)

details_table.cell(0,0).text = "PROGRAM TITLE:"
details_table.cell(0,1).text = program_title

details_table.cell(0,2).text = "DURATION:"
details_table.cell(0,3).text = duration

details_table.cell(1,0).text = "PROGRAM CODE:"
details_table.cell(1,1).text = program_code

details_table.cell(1,2).text = "VENUE:"
details_table.cell(1,3).text = venue

details_table.cell(2,0).text = "COORDINATOR:"
details_table.cell(2,1).text = coordinator

details_table.cell(2,2).text = "PROGRAM ASST:"
details_table.cell(2,3).text = assistant

set_table_borders(details_table)

# =========================================================
# INTRODUCTION
# =========================================================
doc.add_heading(
    "A. PROGRAMME EVALUATION",
    level=2
)

doc.add_paragraph(
    "KSG conducted a programme evaluation to assess the quality, "
    "relevance, and effectiveness of the training. Participants "
    "provided feedback on key aspects of the programme, including "
    "content, delivery, and coordination. The findings will inform "
    "continuous improvement and enhance future programme delivery."
)

# =========================================================
# DETECT COLUMNS
# =========================================================
columns = df.columns.tolist()

objective_col = next(
    (
        col for col in columns
        if "objective" in col.lower()
    ),
    None
)

expectation_col = next(
    (
        col for col in columns
        if "expectation" in col.lower()
    ),
    None
)

future_col = next(
    (
        col for col in columns
        if (
            "attend another online training" in col.lower()
            or "future online training" in col.lower()
            or "attending future" in col.lower()
        )
    ),
    None
)

recommend_col = next(
    (
        col for col in columns
        if (
            "recommend ksg online learning" in col.lower()
            or "colleagues/friends" in col.lower()
            or "recommend" in col.lower()
        )
    ),
    None
)

# =========================================================
# SECTION 1
# =========================================================
doc.add_heading(
    "1. Course Objectives Achievement",
    level=2
)

pct1 = add_percentage_table(

    doc,
    objective_col,

    {
        5:"Excellent",
        4:"Very Good",
        3:"Satisfactory",
        2:"Poor",
        1:"Very Poor"
    }
)

section1_prompt = f"""
Write a short professional interpretation for these results:

Excellent: {pct1[5]}%
Very Good: {pct1[4]}%
Satisfactory: {pct1[3]}%
Poor: {pct1[2]}%
Very Poor: {pct1[1]}%

Focus on achievement of course objectives.
"""

doc.add_paragraph(
    generate_text(section1_prompt)
)

# =========================================================
# SECTION 2
# =========================================================
doc.add_heading(
    "2. Fulfilment of Personal Expectations",
    level=2
)

pct2 = add_percentage_table(

    doc,
    expectation_col,

    {
        5:"5 - Great Extent",
        4:"4 - Some Extent",
        3:"3 - Satisfactory",
        2:"2 - Not Sure",
        1:"1 - Not at All"
    }
)

section2_prompt = f"""
Write a short professional interpretation for these findings:

Great Extent: {pct2[5]}%
Some Extent: {pct2[4]}%
Satisfactory: {pct2[3]}%
Not Sure: {pct2[2]}%
Not at All: {pct2[1]}%

Focus on fulfilment of participant expectations.
"""

doc.add_paragraph(
    generate_text(section2_prompt)
)

# =========================================================
# SECTION 3
# =========================================================
doc.add_heading(
    "3. Please rate the following aspects of the training program:",
    level=2
)

table3 = doc.add_table(
    rows=2,
    cols=6
)

headers1 = [
    "ASPECT OF THE PROGRAM",
    "Excellent %",
    "Very Good %",
    "Satisfactory %",
    "Poor %",
    "Very poor %"
]

headers2 = [
    "",
    "5",
    "4",
    "3",
    "2",
    "1"
]

for i, val in enumerate(headers1):
    table3.rows[0].cells[i].text = val

for i, val in enumerate(headers2):
    table3.rows[1].cells[i].text = val

exclude_cols = [
    objective_col,
    expectation_col,
    future_col,
    recommend_col
]

exclude_keywords = [
    "response",
    "number"
]

rating_cols = []

for col in columns:

    if col not in exclude_cols:

        if pd.api.types.is_numeric_dtype(df[col]):

            if not any(
                keyword in col.lower()
                for keyword in exclude_keywords
            ):

                rating_cols.append(col)

# =========================================================
# CLEAN DISPLAY LABELS
# =========================================================
display_labels = {

    "How do you rate course organisation and co-ordination":
    "Course organisation and coordination",

    "How do you rate content of training programme":
    "Content of training programme",

    "How do you rate relevance of training programme/Course to your Job":
    "Relevance of training programme/Course to your Job",

    "How do you rate quality of training/facilitation and learning materials":
    "Quality of training/facilitation and learning materials",

    "How do you rate the appropriateness of duration of programme(length of course)":
    "Appropriateness of duration of programme (length of course)",

    "How do you rate the appropriateness of online training platform":
    "Appropriateness of online training platform",

    "How do you rate the technical supoport given during the course":
    "Technical support given during the course"
}

for col in rating_cols:

    counts = df[col].value_counts().to_dict()

    total = sum(counts.values())

    excellent = round((counts.get(5,0)/total)*100) if total else 0
    very_good = round((counts.get(4,0)/total)*100) if total else 0
    satisfactory = round((counts.get(3,0)/total)*100) if total else 0
    poor = round((counts.get(2,0)/total)*100) if total else 0
    very_poor = round((counts.get(1,0)/total)*100) if total else 0

    row = table3.add_row().cells

    row[0].text = display_labels.get(col, col)
    row[1].text = str(excellent)
    row[2].text = str(very_good)
    row[3].text = str(satisfactory)
    row[4].text = str(poor)
    row[5].text = str(very_poor)

set_table_borders(table3)

section3_prompt = """
Write one short professional paragraph interpreting the overall ratings.
Highlight strengths and possible improvement areas.
Use concise human institutional language.
"""

doc.add_paragraph(
    generate_text(section3_prompt)
)

# =========================================================
# SECTION 4
# =========================================================
doc.add_heading(
    "4. Likelihood of Attending Future Online Training Sessions at KSG",
    level=2
)

pct4 = add_percentage_table(

    doc,
    future_col,

    {
        5:"5 - Great Extent",
        4:"4 - Some Extent",
        3:"3 - Satisfactory",
        2:"2 - Not Sure",
        1:"1 - Not at All"
    }
)

section4_prompt = f"""
Write a short professional interpretation for:

Great Extent: {pct4[5]}%
Some Extent: {pct4[4]}%
Satisfactory: {pct4[3]}%

Focus on likelihood of future participation.
Use concise human institutional language.
"""

doc.add_paragraph(
    generate_text(section4_prompt)
)

# =========================================================
# SECTION 5
# =========================================================
doc.add_heading(
    "5. Willingness to Recommend KSG Online Learning to Colleagues and Friends",
    level=2
)

pct5 = add_percentage_table(

    doc,
    recommend_col,

    {
        5:"5 - Great Extent",
        4:"4 - Some Extent",
        3:"3 - Satisfactory",
        2:"2 - Not Sure",
        1:"1 - Not at All"
    }
)

section5_prompt = f"""
Write a short professional interpretation for:

Great Extent: {pct5[5]}%
Some Extent: {pct5[4]}%
Satisfactory: {pct5[3]}%

Focus on willingness to recommend KSG online learning.
Use concise human institutional language.
"""

doc.add_paragraph(
    generate_text(section5_prompt)
)

# =========================================================
# QUALITATIVE SECTIONS
# =========================================================
qual_sections = [

    (
        "6. Suggestions for Improving KSG eLearning Courses and Enhancing Participants’ Learning Experiences",
        "improv"
    ),

    (
        "7. Additional Comments and Suggestions on the Training Program",
        "comment"
    ),

    (
        "8. Recommended Additional Topics for Inclusion in the Training Programme",
        "topic"
    ),

    (
        "9. Additional Training Programs of Interest",
        "interest"
    )
]

for title, keyword in qual_sections:

    raw_text = get_qualitative_text(keyword)

    doc.add_heading(title, level=2)

    prompt = f"""
You are an institutional monitoring and evaluation officer at the Kenya School of Government.

The following are participant responses from an online training evaluation:

{raw_text}

Your task:
- Identify the main recurring themes from the responses
- Summarize them briefly in clear human language
- Use a professional but natural KSG reporting tone
- Keep the paragraph concise and evidence-based
- Avoid exaggerated language
- Avoid academic or AI-sounding wording
- Preserve the meaning and participant voice
- Focus only on the most common issues/views raised

Write one short institutional paragraph only.
"""

    doc.add_paragraph(
        generate_text(prompt)
    )

# =========================================================
# SECTION 10
# =========================================================
doc.add_heading(
    "10. Key Insights & Recommendations",
    level=2
)

prompt10 = f"""
Write:
1. Key Insights
2. Recommendations

for an online training evaluation report.

Use these findings:
- Strong ratings on course relevance and facilitation
- Positive learner engagement
- Strong willingness to attend future programs
- Strong willingness to recommend KSG online learning
- Need for improved technical support and platform reliability

Write in concise institutional reporting style using simple human language.
"""

doc.add_paragraph(
    generate_text(prompt10)
)

# =========================================================
# SIGNATURES
# =========================================================
doc.add_paragraph("\n")

doc.add_paragraph(
    "Prepared by: _____________________    "
    "Date: _____________________    "
    "Signature: _____________________"
)

doc.add_paragraph("\n")

doc.add_paragraph(
    "Confirmed by: _____________________    "
    "Date: _____________________    "
    "Signature: _____________________"
)

doc.add_paragraph("\n")

doc.add_paragraph(
    "Approved by: _____________________    "
    "Date: _____________________    "
    "Signature: _____________________"
)

# =========================================================
# SAVE REPORT
# =========================================================
base_name = os.path.splitext(
    os.path.basename(file_name)
)[0]

output_file = f"{base_name}_Online_EEE_Report_LLM.docx"

doc.save(output_file)

print("\n===================================")
print("ONLINE LLM REPORT GENERATED")
print("===================================")

print(f"\nSaved as: {output_file}")
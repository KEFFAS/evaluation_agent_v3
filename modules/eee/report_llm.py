import os
import pandas as pd
from dotenv import load_dotenv
from openai import OpenAI

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ==========================================================
# LOAD API
# ==========================================================
load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)


# ==========================================================
# OPENAI FUNCTION
# ==========================================================
def generate_text(prompt):

    response = client.chat.completions.create(

        model="gpt-5-nano-2025-08-07",

        messages=[

            {
                "role": "system",
                "content": (
                    "You are an institutional monitoring and evaluation "
                    "officer writing formal Kenya School of Government "
                    "evaluation reports. Write in a professional, concise, "
                    "evidence-based and human tone. Avoid exaggerated "
                    "language, repetition, and generic AI wording."
                )
            },

            {
                "role": "user",
                "content": prompt
            }

        ]

    )

    return response.choices[0].message.content.strip()


# ==========================================================
# TABLE BORDERS
# ==========================================================
def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for border_name in [

        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"

    ]:

        border = OxmlElement(f"w:{border_name}")

        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:color"), "000000")

        borders.append(border)

    tblPr.append(borders)


# ==========================================================
# CALCULATE RATING STATISTICS
# ==========================================================
def calculate_rating_stats(series):
    """
    Calculate:

    - Rating counts for 5, 4, 3, 2 and 1
    - Total respondents
    - Mean score
    - Composite score
    """

    numeric = pd.to_numeric(
        series,
        errors="coerce"
    )

    counts = {

        score: int((numeric == score).sum())

        for score in [5, 4, 3, 2, 1]

    }

    total_respondents = sum(
        counts.values()
    )

    if total_respondents > 0:

        mean_score = (

            (5 * counts[5]) +
            (4 * counts[4]) +
            (3 * counts[3]) +
            (2 * counts[2]) +
            (1 * counts[1])

        ) / total_respondents

        composite_score = (
            mean_score / 5
        ) * 100

    else:

        mean_score = 0
        composite_score = 0

    return {

        "count_5": counts[5],
        "count_4": counts[4],
        "count_3": counts[3],
        "count_2": counts[2],
        "count_1": counts[1],

        "total_respondents":
        total_respondents,

        "mean_score":
        round(mean_score, 2),

        "composite_score":
        round(composite_score, 1)

    }


# ==========================================================
# ADD RATING TABLE
# ==========================================================
def add_single_rating_table(
    doc,
    title,
    stats,
    rating_labels
):
    """
    Add FE-style table for one evaluation item.
    """

    doc.add_paragraph(
        f"Total Respondents: "
        f"{stats['total_respondents']}"
    )

    table = doc.add_table(
        rows=1,
        cols=7
    )

    headers = [

        title,
        "5",
        "4",
        "3",
        "2",
        "1",
        "MEAN"

    ]

    for i, header in enumerate(headers):

        table.rows[0].cells[i].text = (
            str(header)
        )

    # ======================================================
    # RATING SCALE
    # ======================================================
    scale_row = table.add_row().cells

    scale_row[0].text = "Rating Scale"

    scale_row[1].text = rating_labels.get(5, "")
    scale_row[2].text = rating_labels.get(4, "")
    scale_row[3].text = rating_labels.get(3, "")
    scale_row[4].text = rating_labels.get(2, "")
    scale_row[5].text = rating_labels.get(1, "")
    scale_row[6].text = ""

    # ======================================================
    # RESULTS
    # ======================================================
    row = table.add_row().cells

    row[0].text = title

    row[1].text = str(
        stats["count_5"]
    )

    row[2].text = str(
        stats["count_4"]
    )

    row[3].text = str(
        stats["count_3"]
    )

    row[4].text = str(
        stats["count_2"]
    )

    row[5].text = str(
        stats["count_1"]
    )

    row[6].text = str(
        stats["mean_score"]
    )

    # ======================================================
    # COMPOSITE SCORE
    # ======================================================
    composite_row = table.add_row().cells

    composite_row[0].text = (
        "COMPOSITE SCORE (%)"
    )

    composite_row[6].text = (
        f"{stats['composite_score']}%"
    )

    set_table_borders(table)

    return table


# ==========================================================
# REPORT FUNCTION
# ==========================================================
def generate_eee_report_llm(

        cleaned_file,

        programme_title,

        programme_code,

        duration,

        venue,

        coordinator,

        assistant,

        output_folder="outputs"

):

    os.makedirs(
        output_folder,
        exist_ok=True
    )

    # =====================================================
    # LOAD DATA
    # =====================================================
    df = pd.read_excel(
        cleaned_file
    )

    print(
        f"Loaded: {cleaned_file}"
    )

    # =====================================================
    # USE PASSED DETAILS
    # =====================================================
    program_title = programme_title
    program_code = programme_code

    # =====================================================
    # DETECT RATING COLUMNS
    # =====================================================
    rating_cols = []

    for col in df.columns:

        numeric = pd.to_numeric(
            df[col],
            errors="coerce"
        )

        valid_values = numeric.dropna()

        if len(valid_values) > 0:

            if valid_values.isin(
                [1, 2, 3, 4, 5]
            ).all():

                rating_cols.append(col)

    # =====================================================
    # EXCLUDE NON-RATING COLUMNS
    # =====================================================
    rating_cols = [

        col

        for col in rating_cols

        if col != "Timetable No"

    ]

    # =====================================================
    # SPECIAL COLUMNS
    # =====================================================
    objective_col = next(

        (

            col

            for col in rating_cols

            if "objective"
            in str(col).lower()

        ),

        None

    )

    expectation_col = next(

        (

            col

            for col in rating_cols

            if "expectation"
            in str(col).lower()

        ),

        None

    )

    comparison_col = next(

        (

            col

            for col in rating_cols

            if (

                "similar institution"
                in str(col).lower()

                or

                "similar institutions"
                in str(col).lower()

            )

        ),

        None

    )

    # =====================================================
    # CALCULATE SECTION STATISTICS
    # =====================================================
    section1_stats = None
    section2_stats = None
    section4_stats = None

    if objective_col:

        section1_stats = (
            calculate_rating_stats(
                df[objective_col]
            )
        )

    if expectation_col:

        section2_stats = (
            calculate_rating_stats(
                df[expectation_col]
            )
        )

    if comparison_col:

        section4_stats = (
            calculate_rating_stats(
                df[comparison_col]
            )
        )

    # =====================================================
    # SPECIFIC PROGRAMME ASPECTS
    # =====================================================
    special_cols = [

        objective_col,
        expectation_col,
        comparison_col

    ]

    specific_aspects = []

    for col in rating_cols:

        if col not in special_cols:

            stats = (
                calculate_rating_stats(
                    df[col]
                )
            )

            specific_aspects.append({

                "aspect": col,

                "stats": stats

            })

    # =====================================================
    # SECTION 3 OVERALL MEAN
    # =====================================================
    if specific_aspects:

        section3_overall_mean = round(

            sum(

                item["stats"]["mean_score"]

                for item
                in specific_aspects

            )

            / len(specific_aspects),

            2

        )

        section3_composite = round(

            (
                section3_overall_mean / 5
            ) * 100,

            1

        )

    else:

        section3_overall_mean = 0
        section3_composite = 0

    # =====================================================
    # OVERALL SUMMARY DATA
    # =====================================================
    summary_data = []

    if section1_stats:

        summary_data.append({

            "area":
            "Course Objectives Achievement",

            "mean":
            section1_stats["mean_score"],

            "composite":
            section1_stats["composite_score"]

        })

    if section2_stats:

        summary_data.append({

            "area":
            "Fulfilment of Personal Expectations",

            "mean":
            section2_stats["mean_score"],

            "composite":
            section2_stats["composite_score"]

        })

    if specific_aspects:

        summary_data.append({

            "area":
            "Specific Programme Aspects",

            "mean":
            section3_overall_mean,

            "composite":
            section3_composite

        })

    if section4_stats:

        summary_data.append({

            "area":
            "KSG Compared to Similar Institutions",

            "mean":
            section4_stats["mean_score"],

            "composite":
            section4_stats["composite_score"]

        })

    # =====================================================
    # OVERALL MEAN
    # =====================================================
    if summary_data:

        overall_mean = round(

            sum(
                item["mean"]
                for item
                in summary_data
            )

            / len(summary_data),

            2

        )

        overall_composite = round(

            (
                overall_mean / 5
            ) * 100,

            1

        )

    else:

        overall_mean = 0
        overall_composite = 0

    # =====================================================
    # DOCUMENT
    # =====================================================
    doc = Document()

    style = doc.styles["Normal"]

    style.font.name = (
        "Times New Roman"
    )

    style.font.size = Pt(11)

    # =====================================================
    # HEADER
    # =====================================================
    doc.add_paragraph(
        "KSG/17/EOEEF/07"
    )

    doc.add_paragraph(
        "KENYA SCHOOL OF GOVERNMENT"
    )

    doc.add_paragraph(
        "MATUGA"
    )

    doc.add_heading(

        "END-OF-EVENT EVALUATION REPORT",

        level=1

    )

    # =====================================================
    # PROGRAM DETAILS
    # =====================================================
    table = doc.add_table(

        rows=3,

        cols=4

    )

    table.cell(
        0, 0
    ).text = "PROGRAMME TITLE:"

    table.cell(
        0, 1
    ).text = str(program_title)

    table.cell(
        0, 2
    ).text = "DURATION:"

    table.cell(
        0, 3
    ).text = str(duration)

    table.cell(
        1, 0
    ).text = "PROGRAM CODE:"

    table.cell(
        1, 1
    ).text = str(program_code)

    table.cell(
        1, 2
    ).text = "VENUE:"

    table.cell(
        1, 3
    ).text = str(venue)

    table.cell(
        2, 0
    ).text = "COORDINATOR:"

    table.cell(
        2, 1
    ).text = str(coordinator)

    table.cell(
        2, 2
    ).text = "PROGRAM ASSISTANT:"

    table.cell(
        2, 3
    ).text = str(assistant)

    set_table_borders(table)

    # =====================================================
    # PROGRAMME EVALUATION
    # =====================================================
    doc.add_heading(

        "A. PROGRAMME EVALUATION",

        level=2

    )

    doc.add_paragraph(

        "KSG conducted a programme evaluation "
        "to assess the quality, relevance and "
        "effectiveness of the training programme. "
        "Participants provided feedback on key "
        "aspects of programme delivery and "
        "administration. The findings will inform "
        "continuous improvement and enhance future "
        "programme delivery."

    )

    # =====================================================
    # OVERALL PROGRAMME SUMMARY
    # =====================================================
    doc.add_heading(

        "1. Overall Programme Evaluation Summary",

        level=2

    )

    summary_table = doc.add_table(

        rows=1,

        cols=3

    )

    summary_headers = [

        "EVALUATION AREA",

        "MEAN SCORE",

        "COMPOSITE SCORE (%)"

    ]

    for i, header in enumerate(
        summary_headers
    ):

        summary_table.rows[0].cells[i].text = (
            header
        )

    for item in summary_data:

        row = (
            summary_table.add_row().cells
        )

        row[0].text = item["area"]

        row[1].text = str(
            item["mean"]
        )

        row[2].text = (
            f"{item['composite']}%"
        )

    # Overall Evaluation Score
    row = summary_table.add_row().cells

    row[0].text = (
        "OVERALL EVALUATION SCORE"
    )

    row[1].text = str(
        overall_mean
    )

    row[2].text = (
        f"{overall_composite}%"
    )

    set_table_borders(summary_table)

    # =====================================================
    # AI SUMMARY INTERPRETATION
    # =====================================================
    summary_prompt = f"""
The following are End-of-Event Evaluation
composite results for a Kenya School of Government
training programme:

{summary_data}

Overall Mean Score:
{overall_mean}

Overall Composite Score:
{overall_composite}%

Write one concise institutional paragraph
interpreting the overall performance.

Requirements:

- Professional and evidence-based.
- Mention the overall level of performance.
- Identify only major strengths or areas requiring
  attention based on the scores.
- Do not exaggerate.
- Do not repeat all the figures.
"""

    doc.add_paragraph(

        generate_text(
            summary_prompt
        )

    )

    # =====================================================
    # SECTION 2
    # COURSE OBJECTIVES
    # =====================================================
    doc.add_heading(

        "2. Course Objectives Achievement",

        level=2

    )

    if section1_stats:

        objective_labels = {

            5: "Excellent",
            4: "Very Good",
            3: "Satisfactory",
            2: "Poor",
            1: "Very Poor"

        }

        add_single_rating_table(

            doc,

            "COURSE OBJECTIVES ACHIEVEMENT",

            section1_stats,

            objective_labels

        )

        prompt = f"""
Interpret the following Course Objectives
Achievement results from a Kenya School of Government
training programme:

5 = {section1_stats['count_5']}
4 = {section1_stats['count_4']}
3 = {section1_stats['count_3']}
2 = {section1_stats['count_2']}
1 = {section1_stats['count_1']}

Mean Score:
{section1_stats['mean_score']}

Composite Score:
{section1_stats['composite_score']}%

Write one concise institutional paragraph.
"""

        doc.add_paragraph(

            generate_text(prompt)

        )

    # =====================================================
    # SECTION 3
    # PERSONAL EXPECTATIONS
    # =====================================================
    doc.add_heading(

        "3. Fulfilment of Personal Expectations",

        level=2

    )

    if section2_stats:

        expectation_labels = {

            5: "Great Extent",
            4: "Some Extent",
            3: "Satisfactory",
            2: "Not Sure",
            1: "Not at All"

        }

        add_single_rating_table(

            doc,

            "FULFILMENT OF EXPECTATIONS",

            section2_stats,

            expectation_labels

        )

        prompt = f"""
Interpret the following Personal Expectations
Fulfilment results from a Kenya School of Government
training programme:

5 = {section2_stats['count_5']}
4 = {section2_stats['count_4']}
3 = {section2_stats['count_3']}
2 = {section2_stats['count_2']}
1 = {section2_stats['count_1']}

Mean Score:
{section2_stats['mean_score']}

Composite Score:
{section2_stats['composite_score']}%

Write one concise institutional paragraph.
"""

        doc.add_paragraph(

            generate_text(prompt)

        )

    # =====================================================
    # SECTION 4
    # SPECIFIC PROGRAMME ASPECTS
    # =====================================================
    doc.add_heading(

        "4. Ratings on Specific Aspects of the Training Programme",

        level=2

    )

    if specific_aspects:

        total_respondents = max(

            item["stats"]["total_respondents"]

            for item
            in specific_aspects

        )

        doc.add_paragraph(

            f"Total Respondents: "
            f"{total_respondents}"

        )

        table3 = doc.add_table(

            rows=1,

            cols=7

        )

        headers = [

            "SPECIFIC ASPECTS",

            "5",

            "4",

            "3",

            "2",

            "1",

            "MEAN"

        ]

        for i, header in enumerate(headers):

            table3.rows[0].cells[i].text = (
                header
            )

        # ================================================
        # RATING SCALE
        # ================================================
        scale_row = (
            table3.add_row().cells
        )

        scale_row[0].text = "Rating Scale"

        scale_row[1].text = "Excellent"
        scale_row[2].text = "Very Good"
        scale_row[3].text = "Satisfactory"
        scale_row[4].text = "Poor"
        scale_row[5].text = "Very Poor"

        # ================================================
        # RESULTS
        # ================================================
        section3_summary = []

        for item in specific_aspects:

            stats = item["stats"]

            row = (
                table3.add_row().cells
            )

            row[0].text = str(
                item["aspect"]
            )

            row[1].text = str(
                stats["count_5"]
            )

            row[2].text = str(
                stats["count_4"]
            )

            row[3].text = str(
                stats["count_3"]
            )

            row[4].text = str(
                stats["count_2"]
            )

            row[5].text = str(
                stats["count_1"]
            )

            row[6].text = str(
                stats["mean_score"]
            )

            section3_summary.append(

                f"{item['aspect']}: "
                f"Mean={stats['mean_score']}, "
                f"Composite="
                f"{stats['composite_score']}%"

            )

        # ================================================
        # OVERALL MEAN
        # ================================================
        overall_row = (
            table3.add_row().cells
        )

        overall_row[0].text = (
            "OVERALL MEAN"
        )

        overall_row[6].text = str(
            section3_overall_mean
        )

        # ================================================
        # COMPOSITE SCORE
        # ================================================
        composite_row = (
            table3.add_row().cells
        )

        composite_row[0].text = (
            "COMPOSITE SCORE (%)"
        )

        composite_row[6].text = (
            f"{section3_composite}%"
        )

        set_table_borders(table3)

        prompt = f"""
Interpret the following Specific Programme
Aspect results from a Kenya School of Government
training evaluation:

{section3_summary}

Overall Mean Score:
{section3_overall_mean}

Overall Composite Score:
{section3_composite}%

Write one concise institutional paragraph.

Highlight the major strengths and any areas requiring
attention based on the evidence.
Do not exaggerate or repeat every score.
"""

        doc.add_paragraph(

            generate_text(prompt)

        )

    # =====================================================
    # QUALITATIVE SECTIONS
    # =====================================================
    qualitative_mapping = {

        "5. Suggestions on the aspects listed in (4) above.":

        "suggestions on aspects",

        "6. Areas to be added to this training programme":

        "other areas you would like added",

        "7. Interest in Other KSG Programmes":

        "other ksg training programs",

        "8. Interest in Additional Training Areas Not Currently Offered by KSG":

        "other training programs not currently offered",

        "10. General Comments":

        "other comments"

    }

    for section_title, keyword in (
        qualitative_mapping.items()
    ):

        matching_cols = [

            col

            for col in df.columns

            if keyword
            in str(col).lower()

        ]

        if matching_cols:

            responses = (

                df[matching_cols[0]]

                .dropna()

                .astype(str)

                .str.strip()

            )

            responses = responses[
                responses != ""
            ]

            responses = (
                responses.drop_duplicates()
            )

            joined_text = (
                " ".join(responses)
            )

            if joined_text.strip():

                prompt = f"""
The following are participant responses from
a Kenya School of Government training evaluation.

Responses:

{joined_text}

Summarize the responses into one concise paragraph.

Requirements:

- Identify only recurring themes.
- Use simple human language.
- Write in Kenya School of Government reporting style.
- Do not list every response.
- Avoid repetition.
- Do not invent information.
"""

                doc.add_heading(

                    section_title,

                    level=2

                )

                doc.add_paragraph(

                    generate_text(prompt)

                )

    # =====================================================
    # SECTION 9
    # INSTITUTIONAL COMPARISON
    # =====================================================
    doc.add_heading(

        "9. Rating of KSG's Training Compared to Similar Institutions",

        level=2

    )

    if section4_stats:

        comparison_labels = {

            5: "Very High",
            4: "High",
            3: "Average",
            2: "Low",
            1: "Very Low"

        }

        add_single_rating_table(

            doc,

            "KSG COMPARED TO SIMILAR INSTITUTIONS",

            section4_stats,

            comparison_labels

        )

        prompt = f"""
Interpret the following institutional comparison
results from a Kenya School of Government
training evaluation:

5 = {section4_stats['count_5']}
4 = {section4_stats['count_4']}
3 = {section4_stats['count_3']}
2 = {section4_stats['count_2']}
1 = {section4_stats['count_1']}

Mean Score:
{section4_stats['mean_score']}

Composite Score:
{section4_stats['composite_score']}%

Write one concise evidence-based paragraph suitable
for a Kenya School of Government evaluation report.
"""

        doc.add_paragraph(

            generate_text(prompt)

        )

    # =====================================================
    # SECTION 11
    # KEY RECOMMENDATIONS
    # =====================================================
    doc.add_heading(

        "11. Key Recommendations",

        level=2

    )

    recommendation_text = ""

    for col in df.columns:

        if any(

            keyword
            in str(col).lower()

            for keyword in [

                "suggest",
                "comment",
                "area"

            ]

        ):

            recommendation_text += " "

            recommendation_text += " ".join(

                df[col]

                .dropna()

                .astype(str)

            )

    prompt = f"""
Participants made the following suggestions
and comments:

{recommendation_text}

Prepare concise institutional recommendations.

Requirements:

- Maximum 5 recommendations.
- Use bullet points.
- Be practical.
- Base recommendations only on participant feedback.
- Use Kenya School of Government reporting style.
- Do not invent recommendations not supported
  by the feedback.
"""

    recommendations = generate_text(
        prompt
    )

    doc.add_paragraph(
        recommendations
    )

    # =====================================================
    # SIGNATURES
    # =====================================================
    doc.add_paragraph()

    doc.add_paragraph(

        "Prepared by............................................................"

    )

    doc.add_paragraph(

        "Date............................   "
        "Signature............................"

    )

    doc.add_paragraph()

    doc.add_paragraph(

        "Confirmed by..........................................................."

    )

    doc.add_paragraph(

        "Date............................   "
        "Signature............................"

    )

    doc.add_paragraph()

    doc.add_paragraph(

        "Approved by............................................................"

    )

    doc.add_paragraph(

        "Date............................   "
        "Signature............................"

    )

    # =====================================================
    # SAVE REPORT
    # =====================================================
    base_name = os.path.splitext(

        os.path.basename(
            cleaned_file
        )

    )[0]

    report_file = os.path.join(

        output_folder,

        f"{base_name}_EEE_Report_LLM.docx"

    )

    doc.save(report_file)

    print(

        f"\nLLM report generated: "
        f"{report_file}"

    )

    return report_file
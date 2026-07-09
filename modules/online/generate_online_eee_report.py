# =========================================================
# ONLINE END-OF-EVENT REPORT GENERATOR
# FINAL INSTITUTIONAL VERSION
# =========================================================

import pandas as pd
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# =========================================================
# TABLE BORDER FUNCTION
# =========================================================
def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement('w:tblBorders')

    for border_name in [
        'top',
        'left',
        'bottom',
        'right',
        'insideH',
        'insideV'
    ]:

        border = OxmlElement(f'w:{border_name}')

        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:color'), '000000')

        borders.append(border)

    tblPr.append(borders)

# =========================================================
# ADD BULLETS
# =========================================================
def add_bullets(items):

    for item in items:

        p = doc.add_paragraph(style='List Bullet')

        run = p.add_run(item)

        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

# =========================================================
# ADD PERCENTAGE TABLE
# =========================================================
def add_percentage_table(doc, column_name, labels):

    table = doc.add_table(rows=1, cols=2)

    table.rows[0].cells[0].text = "Rating"
    table.rows[0].cells[1].text = "Percentage of Respondents"

    counts = df[column_name].value_counts().to_dict()

    total = sum(counts.values())

    percentages = {}

    for score in [5,4,3,2,1]:

        pct = (
            round((counts.get(score,0)/total)*100,1)
            if total > 0 else 0
        )

        percentages[score] = pct

        row = table.add_row().cells

        row[0].text = labels[score]
        row[1].text = str(pct)

    set_table_borders(table)

    return percentages

# =========================================================
# QUALITATIVE TEXTJOIN FUNCTION
# =========================================================
def add_qualitative_textjoin(title, keyword):

    matching_cols = [

        col for col in columns

        if keyword in col.lower()
    ]

    if matching_cols:

        col = matching_cols[0]

        responses = (
            df[col]
            .dropna()
            .astype(str)
            .str.strip()
        )

        responses = responses[
            responses != ""
        ]

        responses = list(dict.fromkeys(responses))

        cleaned_text = " ".join(responses)

        doc.add_heading(
            title,
            level=2
        )

        doc.add_paragraph(cleaned_text)

# =========================================================
# INPUT FILE
# =========================================================
file_name = input(
    "Enter cleaned Online EEE file: "
).strip()

# =========================================================
# LOAD DATA
# =========================================================
df = pd.read_excel(
    file_name,
    engine="openpyxl"
)

print("\nLoaded cleaned dataset successfully.")

# =========================================================
# USER INPUTS
# =========================================================
program_title = input("Enter Program Title: ")
duration = input("Enter Program Duration: ")
program_code = input("Enter Program Code: ")
venue = input("Enter Venue: ")
coordinator = input("Enter Coordinator Name: ")
assistant = input("Enter Program Assistant Name: ")

# =========================================================
# CREATE DOCUMENT
# =========================================================
doc = Document()

style = doc.styles['Normal']

style.font.name = 'Times New Roman'
style.font.size = Pt(11)

# =========================================================
# HEADER
# =========================================================
doc.add_paragraph("KSG/17/EOEEF/07")

p = doc.add_paragraph()

p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

run = p.add_run(
    "KENYA SCHOOL OF GOVERNMENT\n"
    "MATUGA\n\n"
    "END-OF-EVENT EVALUATION REPORT"
)

run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(16)

# =========================================================
# PROGRAM DETAILS TABLE
# =========================================================
details_table = doc.add_table(
    rows=3,
    cols=4
)

details_table.cell(0,0).text = "PROGRAM TITLE:"
details_table.cell(0,1).text = program_title

details_table.cell(0,2).text = "DURATION:"
details_table.cell(0,3).text = duration

details_table.cell(1,0).text = "PROGRAM CODE:"
details_table.cell(1,1).text = program_code

details_table.cell(1,2).text = "VENUE:"
details_table.cell(1,3).text = venue

details_table.cell(2,0).text = "COORDINATOR:"
details_table.cell(2,1).text = coordinator

details_table.cell(2,2).text = "PROGRAM ASST:"
details_table.cell(2,3).text = assistant

set_table_borders(details_table)

# =========================================================
# INTRODUCTION
# =========================================================
doc.add_heading(
    "A. PROGRAMME EVALUATION",
    level=2
)

doc.add_paragraph(
    "KSG conducted a programme evaluation to assess the quality, "
    "relevance, and effectiveness of the online training programme. "
    "Participants provided feedback on key aspects of the programme, "
    "including content, delivery, online learning experience, and "
    "coordination. The findings will support continuous improvement "
    "of future online learning programmes."
)

# =========================================================
# DETECT COLUMNS
# =========================================================
columns = df.columns.tolist()

objective_col = next(
    (
        col for col in columns
        if "objective" in col.lower()
    ),
    None
)

expectation_col = next(
    (
        col for col in columns
        if "expectation" in col.lower()
    ),
    None
)

future_col = next(
    (
        col for col in columns
        if (
            "attend another online training" in col.lower()
            or "future online training" in col.lower()
            or "attending future" in col.lower()
        )
    ),
    None
)

recommend_col = next(
    (
        col for col in columns
        if (
            "recommend ksg online learning" in col.lower()
            or "colleagues/friends" in col.lower()
            or "recommend" in col.lower()
        )
    ),
    None
)

# =========================================================
# SECTION 1
# =========================================================
doc.add_heading(
    "1. Course Objectives Achievement",
    level=2
)

if objective_col:

    pct1 = add_percentage_table(

        doc,
        objective_col,

        {
            5:"Excellent",
            4:"Very Good",
            3:"Satisfactory",
            2:"Poor",
            1:"Very Poor"
        }
    )

    positive = pct1[5] + pct1[4]

    doc.add_paragraph(
        f"The course objectives were successfully achieved, with "
        f"{round(positive,1)}% of participants rating them as "
        f"Excellent or Very Good, indicating a high level of "
        f"effectiveness in meeting the intended training goals."
    )

# =========================================================
# SECTION 2
# =========================================================
doc.add_heading(
    "2. Fulfilment of Personal Expectations",
    level=2
)

if expectation_col:

    pct2 = add_percentage_table(

        doc,
        expectation_col,

        {
            5:"5 - Great Extent",
            4:"4 - Some Extent",
            3:"3 - Satisfactory",
            2:"2 - Not Sure",
            1:"1 - Not at All"
        }
    )

    doc.add_paragraph(
        f"The program largely met participants’ personal expectations, "
        f"with {pct2[5]}% indicating fulfilment to a great extent and "
        f"an overall positive rating, demonstrating strong alignment "
        f"with participants’ learning needs and expectations."
    )

# =========================================================
# SECTION 3
# =========================================================
doc.add_heading(
    "3. Please rate the following aspects of the training program:",
    level=2
)

table3 = doc.add_table(
    rows=2,
    cols=6
)

headers1 = [
    "ASPECT OF THE PROGRAM",
    "Excellent %",
    "Very Good %",
    "Satisfactory %",
    "Poor %",
    "Very poor %"
]

headers2 = [
    "",
    "5",
    "4",
    "3",
    "2",
    "1"
]

for i, val in enumerate(headers1):
    table3.rows[0].cells[i].text = val

for i, val in enumerate(headers2):
    table3.rows[1].cells[i].text = val

exclude_cols = [
    objective_col,
    expectation_col,
    future_col,
    recommend_col
]

exclude_keywords = [
    "response",
    "number"
]

rating_cols = []

for col in columns:

    if col not in exclude_cols:

        if pd.api.types.is_numeric_dtype(df[col]):

            if not any(
                keyword in col.lower()
                for keyword in exclude_keywords
            ):

                rating_cols.append(col)

# =========================================================
# CLEAN DISPLAY LABELS
# =========================================================
display_labels = {

    "How do you rate course organisation and co-ordination":
    "Course organisation and coordination",

    "How do you rate content of training programme":
    "Content of training programme",

    "How do you rate relevance of training programme/Course to your Job":
    "Relevance of training programme/Course to your Job",

    "How do you rate quality of training/facilitation and learning materials":
    "Quality of training/facilitation and learning materials",

    "How do you rate the appropriateness of duration of programme(length of course)":
    "Appropriateness of duration of programme (length of course)",

    "How do you rate the appropriateness of online training platform":
    "Appropriateness of online training platform",

    "How do you rate the technical supoport given during the course":
    "Technical support given during the course"
}

for col in rating_cols:

    counts = df[col].value_counts().to_dict()

    total = sum(counts.values())

    excellent = round((counts.get(5,0)/total)*100) if total else 0
    very_good = round((counts.get(4,0)/total)*100) if total else 0
    satisfactory = round((counts.get(3,0)/total)*100) if total else 0
    poor = round((counts.get(2,0)/total)*100) if total else 0
    very_poor = round((counts.get(1,0)/total)*100) if total else 0

    row = table3.add_row().cells

    row[0].text = display_labels.get(col, col)
    row[1].text = str(excellent)
    row[2].text = str(very_good)
    row[3].text = str(satisfactory)
    row[4].text = str(poor)
    row[5].text = str(very_poor)

set_table_borders(table3)

doc.add_paragraph(
    "The programme was highly rated across key aspects including "
    "organization, content relevance, facilitation, platform usability, "
    "and technical support."
)

# =========================================================
# SECTION 4
# =========================================================
doc.add_heading(
    "4. Likelihood of Attending Future Online Training Sessions at KSG",
    level=2
)

if future_col:

    pct4 = add_percentage_table(

        doc,
        future_col,

        {
            5:"5 - Great Extent",
            4:"4 - Some Extent",
            3:"3 - Satisfactory",
            2:"2 - Not Sure",
            1:"1 - Not at All"
        }
    )

    positive4 = pct4[5] + pct4[4]

    doc.add_paragraph(
        f"There is a strong likelihood of continued engagement with "
        f"KSG online training programmes, with {round(positive4,1)}% "
        f"of participants expressing willingness to attend future "
        f"sessions to a great or some extent."
    )

# =========================================================
# SECTION 5
# =========================================================
doc.add_heading(
    "5. Willingness to Recommend KSG Online Learning to Colleagues and Friends",
    level=2
)

if recommend_col:

    pct5 = add_percentage_table(

        doc,
        recommend_col,

        {
            5:"5 - Great Extent",
            4:"4 - Some Extent",
            3:"3 - Satisfactory",
            2:"2 - Not Sure",
            1:"1 - Not at All"
        }
    )

    positive5 = pct5[5] + pct5[4]

    doc.add_paragraph(
        f"Participants demonstrated a strong willingness to recommend "
        f"KSG online learning, with {round(positive5,1)}% indicating "
        f"they would do so to a great or some extent, reflecting high "
        f"satisfaction and confidence in the programme."
    )

# =========================================================
# SECTION 6
# =========================================================
add_qualitative_textjoin(
    "6. Suggestions for Improving KSG eLearning Courses and Enhancing Participants’ Learning Experiences",
    "improv"
)

# =========================================================
# SECTION 7
# =========================================================
add_qualitative_textjoin(
    "7. Additional Comments and Suggestions on the Training Program",
    "comment"
)

# =========================================================
# SECTION 8
# =========================================================
add_qualitative_textjoin(
    "8. Recommended Additional Topics for Inclusion in the Training Programme",
    "topic"
)

# =========================================================
# SECTION 9
# =========================================================
add_qualitative_textjoin(
    "9. Additional Training Programs of Interest",
    "interest"
)

# =========================================================
# SECTION 10
# =========================================================
doc.add_heading(
    "10. Key Insights & Recommendations",
    level=2
)

doc.add_heading(
    "Key Insights",
    level=3
)

insights = [
    "The program was highly effective, with the majority of participants rating course objectives achievement positively.",
    "The training content was highly relevant to participants’ workplace responsibilities.",
    "Participants demonstrated strong willingness to continue engaging with KSG online learning programmes.",
    "The programme showed strong coordination, facilitation, and quality of learning materials.",
    "Participants highlighted challenges related to platform reliability and technical support.",
    "There is need to enhance learner engagement through more interactive online learning approaches."
]

add_bullets(insights)

doc.add_heading(
    "Recommendations",
    level=3
)

recommendations = [
    "Improve the reliability and accessibility of the e-learning platform.",
    "Strengthen technical support and participant onboarding mechanisms.",
    "Increase learner engagement through interactive online learning approaches.",
    "Adopt flexible scheduling and provide recorded learning sessions.",
    "Enhance course design using practical and multimedia learning approaches."
]

add_bullets(recommendations)

# =========================================================
# SIGNATURES
# =========================================================
doc.add_paragraph("\n")

doc.add_paragraph(
    "Prepared by: _____________________    "
    "Date: _____________________    "
    "Signature: _____________________"
)

doc.add_paragraph("\n")

doc.add_paragraph(
    "Confirmed by: _____________________    "
    "Date: _____________________    "
    "Signature: _____________________"
)

doc.add_paragraph("\n")

doc.add_paragraph(
    "Approved by: _____________________    "
    "Date: _____________________    "
    "Signature: _____________________"
)

# =========================================================
# SAVE REPORT
# =========================================================
base_name = os.path.splitext(
    os.path.basename(file_name)
)[0]

output_file = f"{base_name}_Online_EEE_Report.docx"

doc.save(output_file)

# =========================================================
# COMPLETE
# =========================================================
print("\n===================================")
print("ONLINE EEE REPORT GENERATED")
print("===================================")

print(f"\nSaved as:")
print(output_file)
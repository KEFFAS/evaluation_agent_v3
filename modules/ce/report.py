import os
import re
import pandas as pd

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from openai import OpenAI
from dotenv import load_dotenv


# =========================================================
# LOAD ENVIRONMENT
# =========================================================

load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)


# =========================================================
# TABLE BORDERS
# =========================================================

def set_table_borders(table):

    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for border_name in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ]:

        border = OxmlElement(f"w:{border_name}")

        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")

        borders.append(border)

    tblPr.append(borders)


# =========================================================
# QUALITATIVE ANALYSIS
# =========================================================

def analyze_qualitative(text):

    if not text or not text.strip():

        return (
            "Participants did not provide sufficient qualitative feedback "
            "for detailed analysis.\n\n"
            "Participants did not provide sufficient suggestions for "
            "improvement."
        )

    # Prevent excessively large prompts
    text = text[:5000]

    prompt = f"""
The following comments were provided by participants during the
evaluation of a Programme Coordinator at the Kenya School of Government.

Participant Feedback:

{text}

Write exactly TWO concise professional paragraphs.

Paragraph 1:
Summarize the most recurring positive feedback regarding programme
coordination, organization, communication, administration,
participant support, responsiveness and overall programme management.

Paragraph 2:
Summarize the most recurring suggestions for improvement. Focus only
on issues mentioned repeatedly or issues that appear significant.

Requirements:
- Use formal institutional language
- Sound natural and human
- Be evidence-based
- Avoid exaggerated praise
- Avoid generic AI wording
- Avoid repetition
- Do not use bullet points
- Do not use headings
- Write a flowing narrative suitable for an official evaluation report
"""

    try:

        response = client.chat.completions.create(

            model="gpt-5-nano-2025-08-07",

            messages=[

                {
                    "role": "system",
                    "content": (
                        "You are an institutional monitoring and evaluation "
                        "officer writing formal Kenya School of Government "
                        "evaluation reports. Write professionally, concisely "
                        "and in a natural evidence-based tone."
                    )
                },

                {
                    "role": "user",
                    "content": prompt
                }

            ]

        )

        return (
            response
            .choices[0]
            .message
            .content
            .strip()
        )

    except Exception as e:

        print(f"LLM qualitative analysis error: {e}")

        return (
            "Participants generally expressed positive views regarding "
            "programme coordination, organization and administration.\n\n"
            "The feedback should be considered in strengthening areas "
            "that require further improvement."
        )


# =========================================================
# GENERATE HOD FEEDBACK USING LLM
# =========================================================

def generate_hod_feedback(
    composite_score,
    overall_mean,
    highest_indicator,
    highest_score,
    lowest_indicator,
    lowest_score
):

    prompt = f"""
You are assisting the Head of Department – Training at the Kenya
School of Government to provide feedback on a Programme Coordinator
Evaluation Report.

Evaluation Results:

Overall Mean Score: {overall_mean} out of 5
Composite Score: {composite_score}%

Highest Rated Aspect:
{highest_indicator}
Mean Score: {highest_score}

Lowest Rated Aspect:
{lowest_indicator}
Mean Score: {lowest_score}

Generate exactly TWO professional sentences.

Sentence 1:
Write a concise Head of Department comment interpreting the
Coordinator's overall performance based on the evaluation results.

Sentence 2:
Write a concise and practical Head of Department recommendation
based on the evaluation results.

Requirements:
- Exactly TWO sentences
- Sentence 1 is the HOD comment
- Sentence 2 is the HOD recommendation
- Formal institutional language
- Natural and human sounding
- Evidence-based
- Avoid exaggerated praise
- Avoid generic wording
- Vary the wording naturally
- Do not use headings
- Do not use bullet points
- Do not number the sentences
"""

    try:

        response = client.chat.completions.create(

            model="gpt-5-nano-2025-08-07",

            messages=[

                {
                    "role": "system",
                    "content": (
                        "You are a senior monitoring and evaluation officer "
                        "supporting the Head of Department – Training at the "
                        "Kenya School of Government. Write concise, natural, "
                        "professional and evidence-based performance feedback."
                    )
                },

                {
                    "role": "user",
                    "content": prompt
                }

            ]

        )

        feedback = (
            response
            .choices[0]
            .message
            .content
            .strip()
        )

        # Clean extra spaces
        feedback = re.sub(
            r"\s+",
            " ",
            feedback
        ).strip()

        # Split into sentences
        sentences = re.split(
            r"(?<=[.!?])\s+",
            feedback
        )

        sentences = [

            sentence.strip()

            for sentence in sentences

            if sentence.strip()

        ]

        # HOD COMMENT
        if len(sentences) >= 1:

            hod_comment = sentences[0]

        else:

            hod_comment = (
                f"The Coordinator attained a composite score of "
                f"{composite_score}%, with {highest_indicator} "
                f"emerging as a key area of strength."
            )

        # HOD RECOMMENDATION
        if len(sentences) >= 2:

            hod_recommendation = sentences[1]

        else:

            hod_recommendation = (
                f"Attention should be directed towards strengthening "
                f"{lowest_indicator} to improve overall programme "
                f"coordination."
            )

        return hod_comment, hod_recommendation

    except Exception as e:

        print(f"LLM HOD feedback error: {e}")

        hod_comment = (
            f"The Coordinator attained a composite score of "
            f"{composite_score}%, with {highest_indicator} "
            f"emerging as a key area of strength."
        )

        hod_recommendation = (
            f"Attention should be directed towards strengthening "
            f"{lowest_indicator} to improve overall programme "
            f"coordination."
        )

        return hod_comment, hod_recommendation


# =========================================================
# MAIN REPORT FUNCTION
# =========================================================

def generate_ce_report(

    cleaned_file,

    programme_title,

    programme_code,

    duration,

    venue,

    coordinator,

    assistant,

    output_folder="outputs"

):

    # =====================================================
    # CREATE OUTPUT FOLDER
    # =====================================================

    os.makedirs(
        output_folder,
        exist_ok=True
    )

    # =====================================================
    # LOAD DATA
    # =====================================================

    df = pd.read_excel(
        cleaned_file
    )

    print(
        "Generating Coordinator Evaluation Report..."
    )

    # =====================================================
    # GET COORDINATOR NAME
    # =====================================================

    if (
        "Coordinator Name" in df.columns
        and not df["Coordinator Name"].dropna().empty
    ):

        coordinator_name = (
            df["Coordinator Name"]
            .dropna()
            .iloc[0]
        )

    else:

        coordinator_name = coordinator

    # =====================================================
    # RATING COLUMNS
    # =====================================================

    expected_rating_cols = [

        "Organization Of Program Opening And Closing",

        "Briefing Participants And Orientation",

        "Leveling Of Participant Expectations",

        "Communication And Feedback",

        "Management Of Timetable And Facilitators",

        "Monitoring Participants Attendance",

        "Program Evaluation",

        "Action Planning",

        "General Administration Of Program"

    ]

    rating_cols = [

        col

        for col in expected_rating_cols

        if col in df.columns

    ]

    if not rating_cols:

        raise ValueError(
            "No Coordinator Evaluation rating columns were found."
        )

    # =====================================================
    # ANALYZE RATINGS
    # =====================================================

    rating_results = []

    indicator_means = {}

    respondent_counts = []

    for col in rating_cols:

        ratings = pd.to_numeric(
            df[col],
            errors="coerce"
        )

        valid_ratings = ratings[
            ratings.isin([1, 2, 3, 4, 5])
        ]

        # Actual respondent counts
        count5 = int(
            (valid_ratings == 5).sum()
        )

        count4 = int(
            (valid_ratings == 4).sum()
        )

        count3 = int(
            (valid_ratings == 3).sum()
        )

        count2 = int(
            (valid_ratings == 2).sum()
        )

        count1 = int(
            (valid_ratings == 1).sum()
        )

        total_valid = len(
            valid_ratings
        )

        respondent_counts.append(
            total_valid
        )

        # Mean score
        if total_valid > 0:

            mean_score = round(
                valid_ratings.mean(),
                2
            )

        else:

            mean_score = 0

        indicator_means[col] = mean_score

        rating_results.append({

            "indicator": col,

            "count5": count5,

            "count4": count4,

            "count3": count3,

            "count2": count2,

            "count1": count1,

            "mean": mean_score

        })

    # =====================================================
    # TOTAL RESPONDENTS
    # =====================================================

    total_respondents = (

        max(respondent_counts)

        if respondent_counts

        else 0

    )

    # =====================================================
    # OVERALL MEAN
    # =====================================================

    valid_means = [

        score

        for score in indicator_means.values()

        if score > 0

    ]

    if valid_means:

        overall_mean = round(

            sum(valid_means) /
            len(valid_means),

            2

        )

    else:

        overall_mean = 0

    # =====================================================
    # COMPOSITE SCORE
    # =====================================================

    composite_score = round(

        (overall_mean / 5) * 100,

        1

    ) if overall_mean > 0 else 0

    # =====================================================
    # HIGHEST AND LOWEST INDICATORS
    # =====================================================

    if valid_means:

        highest_indicator = max(

            indicator_means,

            key=indicator_means.get

        )

        lowest_indicator = min(

            indicator_means,

            key=indicator_means.get

        )

        highest_score = indicator_means[
            highest_indicator
        ]

        lowest_score = indicator_means[
            lowest_indicator
        ]

    else:

        highest_indicator = "N/A"

        highest_score = 0

        lowest_indicator = "N/A"

        lowest_score = 0

    # =====================================================
    # GENERATE HOD FEEDBACK
    # =====================================================

    hod_comment, hod_recommendation = (

        generate_hod_feedback(

            composite_score,

            overall_mean,

            highest_indicator,

            highest_score,

            lowest_indicator,

            lowest_score

        )

    )

    # =====================================================
    # QUALITATIVE DATA
    # =====================================================

    likes = (

        "; ".join(

            df["Like"]
            .dropna()
            .astype(str)

        )

        if "Like" in df.columns

        else ""

    )

    suggestions = (

        "; ".join(

            df["Suggestions"]
            .dropna()
            .astype(str)

        )

        if "Suggestions" in df.columns

        else ""

    )

    combined_text = f"""
Most Liked:
{likes}

Suggestions:
{suggestions}
"""

    qualitative = analyze_qualitative(
        combined_text
    )

    paragraphs = [

        paragraph.strip()

        for paragraph in qualitative.split("\n\n")

        if paragraph.strip()

    ]

    # =====================================================
    # CREATE WORD DOCUMENT
    # =====================================================

    doc = Document()

    style = doc.styles["Normal"]

    style.font.name = "Times New Roman"

    style.font.size = Pt(11)

    # =====================================================
    # REPORT HEADER
    # =====================================================

    doc.add_paragraph(
        "KENYA SCHOOL OF GOVERNMENT"
    )

    doc.add_paragraph(
        "MATUGA CAMPUS"
    )

    doc.add_heading(
        "PROGRAMME COORDINATOR EVALUATION REPORT",
        level=1
    )

    # =====================================================
    # PROGRAMME DETAILS
    # =====================================================

    details_table = doc.add_table(
        rows=4,
        cols=2
    )

    details_table.cell(
        0,
        0
    ).text = "PROGRAMME TITLE:"

    details_table.cell(
        0,
        1
    ).text = str(
        programme_title
    )

    details_table.cell(
        1,
        0
    ).text = "PROGRAMME CODE:"

    details_table.cell(
        1,
        1
    ).text = str(
        programme_code
    )

    details_table.cell(
        2,
        0
    ).text = "DURATION:"

    details_table.cell(
        2,
        1
    ).text = str(
        duration
    )

    details_table.cell(
        3,
        0
    ).text = "COORDINATOR:"

    details_table.cell(
        3,
        1
    ).text = str(
        coordinator_name
    )

    # Bold labels
    for row in details_table.rows:

        for run in row.cells[0].paragraphs[0].runs:

            run.bold = True

    set_table_borders(
        details_table
    )

    # =====================================================
    # INTRODUCTION
    # =====================================================

    doc.add_paragraph()

    doc.add_paragraph(

        "This report provides information to Kenya School of Government "
        "management for decision making and continuous improvement. "
        "Programme Coordinator Evaluation forms are completed by "
        "participants during training programmes to assess programme "
        "coordination, organization, administration and participant support."

    )

    # =====================================================
    # TOTAL RESPONDENTS
    # =====================================================

    doc.add_paragraph()

    respondents_paragraph = doc.add_paragraph()

    respondents_paragraph.add_run(
        "Total Respondents: "
    ).bold = True

    respondents_paragraph.add_run(
        str(total_respondents)
    )

    # =====================================================
    # PARTICIPANTS RATINGS
    # =====================================================

    doc.add_paragraph()

    doc.add_paragraph(
        "I. Participants' Ratings"
    )

    doc.add_paragraph(

        "Rating Scale: 5 = Excellent | 4 = Very Good | "
        "3 = Good | 2 = Fair | 1 = Poor"

    )

    # =====================================================
    # RATINGS TABLE
    # =====================================================

    table = doc.add_table(
        rows=1,
        cols=7
    )

    headers = [

        "SPECIFIC ASPECTS",

        "EXCELLENT\n5",

        "VERY GOOD\n4",

        "GOOD\n3",

        "FAIR\n2",

        "POOR\n1",

        "MEAN"

    ]

    # Add headers
    for i, header in enumerate(headers):

        cell = table.rows[0].cells[i]

        cell.text = header

        for run in cell.paragraphs[0].runs:

            run.bold = True

    # =====================================================
    # ADD RATING RESULTS
    # =====================================================

    for result in rating_results:

        row = table.add_row().cells

        row[0].text = str(
            result["indicator"]
        )

        row[1].text = str(
            result["count5"]
        )

        row[2].text = str(
            result["count4"]
        )

        row[3].text = str(
            result["count3"]
        )

        row[4].text = str(
            result["count2"]
        )

        row[5].text = str(
            result["count1"]
        )

        row[6].text = str(
            result["mean"]
        )

    # =====================================================
    # OVERALL MEAN ROW
    # =====================================================

    mean_row = table.add_row().cells

    mean_row[0].text = (
        "OVERALL MEAN"
    )

    for i in range(1, 6):

        mean_row[i].text = ""

    mean_row[6].text = str(
        overall_mean
    )

    for cell in mean_row:

        for run in cell.paragraphs[0].runs:

            run.bold = True

    # =====================================================
    # COMPOSITE SCORE ROW
    # =====================================================

    composite_row = table.add_row().cells

    composite_row[0].text = (
        "COMPOSITE SCORE (%)"
    )

    for i in range(1, 6):

        composite_row[i].text = ""

    composite_row[6].text = (
        f"{composite_score}%"
    )

    for cell in composite_row:

        for run in cell.paragraphs[0].runs:

            run.bold = True

    set_table_borders(
        table
    )

    # =====================================================
    # PERFORMANCE SUMMARY
    # =====================================================

    doc.add_paragraph()

    summary = doc.add_paragraph()

    summary.add_run(
        "Performance Summary: "
    ).bold = True

    summary.add_run(

        f"The Programme Coordinator attained an overall mean score of "
        f"{overall_mean} out of 5, translating to a composite score of "
        f"{composite_score}%. The highest-rated aspect was "
        f"{highest_indicator} with a mean score of {highest_score}, while "
        f"{lowest_indicator} recorded the lowest mean score of "
        f"{lowest_score}."

    )

    # =====================================================
    # MOST LIKED
    # =====================================================

    doc.add_paragraph()

    doc.add_paragraph(

        "II. Most liked about programme coordination and administration"

    )

    doc.add_paragraph(

        paragraphs[0]

        if len(paragraphs) > 0

        else "No comments were provided."

    )

    # =====================================================
    # SUGGESTIONS
    # =====================================================

    doc.add_paragraph()

    doc.add_paragraph(

        "III. Suggestions on areas of improvement"

    )

    doc.add_paragraph(

        paragraphs[1]

        if len(paragraphs) > 1

        else (
            "Participants expressed minimal suggestions for improvement."
        )

    )

    # =====================================================
    # HOD COMMENTS
    # =====================================================

    doc.add_paragraph()

    doc.add_paragraph(

        "IV. Head of Department – Training's comments:"

    )

    doc.add_paragraph(
        hod_comment
    )

    # =====================================================
    # HOD RECOMMENDATION
    # =====================================================

    doc.add_paragraph()

    doc.add_paragraph(

        "V. Head of Department – Training's proposals or recommendations:"

    )

    doc.add_paragraph(
        hod_recommendation
    )

    # =====================================================
    # HOD SIGNATURE
    # =====================================================

    doc.add_paragraph()

    doc.add_paragraph(
        "Head of Department – Training"
    )

    doc.add_paragraph(
        "Signature: ______________________________"
    )

    doc.add_paragraph(
        "Date: __________________________________"
    )

    # =====================================================
    # SAVE REPORT
    # =====================================================

    base_name = os.path.splitext(

        os.path.basename(cleaned_file)

    )[0]

    output_file = os.path.join(

        output_folder,

        f"{base_name}_CE_Report.docx"

    )

    doc.save(
        output_file
    )

    print()
    print("=" * 60)
    print("COORDINATOR EVALUATION REPORT GENERATED")
    print("=" * 60)
    print(f"Total Respondents : {total_respondents}")
    print(f"Overall Mean      : {overall_mean}")
    print(f"Composite Score   : {composite_score}%")
    print(f"Highest Aspect    : {highest_indicator}")
    print(f"Lowest Aspect     : {lowest_indicator}")
    print(f"Report File       : {output_file}")

    return output_file